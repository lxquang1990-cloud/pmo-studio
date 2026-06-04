from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any
import json

from openpyxl import load_workbook

BAD_NO_SNAPSHOT = "Chưa có snapshot dữ liệu realtime độc lập"

@dataclass(frozen=True)
class QualityIssue:
    severity: str
    code: str
    message: str

@dataclass(frozen=True)
class QualityResult:
    passed: bool
    issues: list[QualityIssue]
    counts: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return {"passed": self.passed, "counts": self.counts, "issues": [i.__dict__ for i in self.issues]}


def check_oracle_workbook(path: str | Path, min_cases: int = 1, require_oracle_sheet: bool = True) -> QualityResult:
    p = Path(path)
    wb = load_workbook(p, data_only=True)
    issues: list[QualityIssue] = []
    if require_oracle_sheet and "03_Oracle_Snapshot" not in wb.sheetnames:
        issues.append(QualityIssue("error", "MISSING_ORACLE_SHEET", "Workbook thiếu sheet 03_Oracle_Snapshot."))
    if "02_TatCa_TestCases" not in wb.sheetnames:
        issues.append(QualityIssue("error", "MISSING_TESTCASE_SHEET", "Workbook thiếu sheet 02_TatCa_TestCases."))
        return QualityResult(False, issues, {"test_cases": 0})
    ws = wb["02_TatCa_TestCases"]
    header_row = 3
    headers = {str(ws.cell(header_row, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    required = ["Mã TC", "Dữ liệu test", "Kết quả mong đợi", "Trạng thái", "Ghi chú"]
    for h in required:
        if h not in headers:
            issues.append(QualityIssue("error", "MISSING_COLUMN", f"Thiếu cột bắt buộc: {h}"))
    test_cases = 0
    no_snapshot_hits = 0
    missing_oracle_ref = 0
    status_counts: dict[str, int] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        tc = ws.cell(r, headers.get("Mã TC", 1)).value
        if not tc:
            continue
        test_cases += 1
        row_text = "\n".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))
        if BAD_NO_SNAPSHOT in row_text:
            no_snapshot_hits += 1
        if "Oracle" not in row_text and str(tc).startswith("ORC-"):
            missing_oracle_ref += 1
        st_col = headers.get("Trạng thái")
        if st_col:
            st = str(ws.cell(r, st_col).value or "").strip() or "<blank>"
            status_counts[st] = status_counts.get(st, 0) + 1
    if test_cases < min_cases:
        issues.append(QualityIssue("error", "TOO_FEW_CASES", f"Workbook chỉ có {test_cases} testcase, nhỏ hơn ngưỡng {min_cases}."))
    if no_snapshot_hits:
        issues.append(QualityIssue("error", "STALE_NO_SNAPSHOT_TEXT", f"Còn {no_snapshot_hits} dòng chứa câu '{BAD_NO_SNAPSHOT}' dù pipeline oracle yêu cầu snapshot."))
    if missing_oracle_ref:
        issues.append(QualityIssue("warning", "MISSING_ORACLE_REF", f"Có {missing_oracle_ref} testcase ORC thiếu Oracle reference trong nội dung."))
    counts = {"test_cases": test_cases, "status": status_counts, "no_snapshot_hits": no_snapshot_hits}
    return QualityResult(not any(i.severity == "error" for i in issues), issues, counts)


def write_quality_report(result: QualityResult, out: str | Path) -> Path:
    p = Path(out)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(result.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
    return p

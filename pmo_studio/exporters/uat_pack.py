"""Test Case Excel + UAT Pack exporter (v2.8)."""
from __future__ import annotations
import re, zipfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def _ids(text, prefix): return sorted(set(re.findall(rf'\b({prefix}-[A-Z0-9-]*\d+)\b', text)))

def export_test_case_excel(project_root: Path) -> Path:
    text='\n'.join(p.read_text(encoding='utf-8', errors='replace') for p in project_root.glob('artifacts/ba/**/*.md') if p.is_file())
    reqs=_ids(text,'REQ'); acs=_ids(text,'AC'); tcs=_ids(text,'TC') or [f'TC-{i:03d}' for i in range(1, max(2,len(acs))+1)]
    outdir=project_root/'exports'/'customer'/'uat-pack'; outdir.mkdir(parents=True, exist_ok=True)
    out=outdir/'test-cases-uat.xlsx'; wb=Workbook(); ws=wb.active; ws.title='Test Cases'
    headers=['TC ID','Linked REQ','Linked AC','Scenario','Precondition','Steps','Expected Result','Actual Result','Status','Evidence','Owner']
    ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid', fgColor='305496'); c.alignment=Alignment(horizontal='center')
    for i,tc in enumerate(tcs,1): ws.append([tc, reqs[(i-1)%len(reqs)] if reqs else '', acs[(i-1)%len(acs)] if acs else '', f'Validate scenario {i}', 'User has required role and test data', 'Execute business workflow and capture evidence', 'System returns expected result and audit/trace is recorded', '', 'Not Run', '', 'Customer UAT'])
    ws2=wb.create_sheet('UAT Signoff'); ws2.append(['Item','Value']);
    for row in [('Project',project_root.name),('UAT Window','TBD by customer'),('Entry Criteria','All high findings closed'),('Exit Criteria','All P0/P1 test cases Passed or accepted with waiver'),('Signoff','Pending')]: ws2.append(row)
    ws3=wb.create_sheet('Traceability'); ws3.append(['REQ','AC','TC']);
    for i,tc in enumerate(tcs,1): ws3.append([reqs[(i-1)%len(reqs)] if reqs else '', acs[(i-1)%len(acs)] if acs else '', tc])
    for wsx in wb.worksheets:
        for col in wsx.columns: wsx.column_dimensions[col[0].column_letter].width=max(12,min(45,max(len(str(c.value or '')) for c in col)+2))
    wb.save(out); return out

def export_uat_pack(project_root: Path) -> Path:
    xlsx=export_test_case_excel(project_root); out=xlsx.parent/'uat-pack.zip'
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, xlsx.name)
        for rel in ['artifacts/ba/05-test-cases.md','artifacts/ic/04-uat-plan.md','quality/customer-review.md','quality/intelligence.md']:
            p=project_root/rel
            if p.exists(): z.write(p, rel)
    return out

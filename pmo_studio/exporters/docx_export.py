"""DOCX exporter for client/internal profiles."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None


def _is_md_table_line(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def _parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if cells and all(set(c) <= {"-", ":", " "} for c in cells):
            continue
        rows.append(cells)
    return rows


def _add_table(document, table_lines: list[str]) -> None:
    rows = _parse_table(table_lines)
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(width):
            table.cell(r_idx, c_idx).text = row[c_idx] if c_idx < len(row) else ""


def _flush_table(document, table_buffer: list[str]) -> None:
    if table_buffer:
        _add_table(document, table_buffer)
        table_buffer.clear()


def _add_markdown(document, text: str) -> None:
    table_buffer: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if _is_md_table_line(stripped):
            table_buffer.append(stripped)
            continue
        _flush_table(document, table_buffer)
        if stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped:
            document.add_paragraph(stripped)
    _flush_table(document, table_buffer)


def export_docx(project_root: Path, profile: str = "client-ready") -> Path:
    if Document is None:
        raise RuntimeError("python-docx is not installed")
    doc = Document()
    meta = _project_meta(project_root)
    doc.add_heading("PMO Studio Documentation Pack", 0)
    doc.add_paragraph(f"Project: {meta.get('project_slug', project_root.name)}")
    doc.add_paragraph(f"Customer: {meta.get('customer', 'N/A')}")
    doc.add_paragraph(f"Profile: {profile}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_heading("Document Index", level=1)
    order = _profile_order(profile)
    for idx, rel in enumerate(order, start=1):
        if (project_root / rel).exists():
            doc.add_paragraph(f"{idx}. {rel}", style="List Number")
    doc.add_heading("Profile Notes", level=1)
    doc.add_paragraph(_profile_note(profile))
    for rel in order:
        path = project_root / rel
        if path.is_dir():
            for child in sorted(path.rglob("*.md")):
                doc.add_page_break()
                doc.add_paragraph(f"Source artifact: {child.relative_to(project_root)}")
                _add_markdown(doc, child.read_text(encoding="utf-8", errors="ignore"))
        elif path.exists():
            doc.add_page_break()
            doc.add_paragraph(f"Source artifact: {rel}")
            _add_markdown(doc, path.read_text(encoding="utf-8", errors="ignore"))
    out_dir = project_root / "exports" / profile
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "pmo-documentation-pack.docx"
    doc.save(out)
    return out


def _project_meta(project_root: Path) -> dict:
    try:
        return json.loads((project_root / "config.json").read_text(encoding="utf-8"))
    except Exception:
        return {}


def _profile_order(profile: str) -> list[str]:
    common_ba = [
        "artifacts/stage-0/project-brief.md",
        "artifacts/ba/01-prd.md",
        "artifacts/ba/02-brd.md",
        "artifacts/ba/03-srs/srs.md",
        "artifacts/ba/05-test-cases.md",
        "traceability/rtm.md",
    ]
    if profile == "client-ready":
        return common_ba + ["artifacts/pm/01-charter.md", "artifacts/ic/04-uat-plan.md"]
    if profile == "developer":
        return common_ba + ["artifacts/ba/03-srs/apis", "artifacts/ba/03-srs/workflows", "artifacts/ic/03-deployment-plan.md"]
    if profile == "management":
        return ["artifacts/po/01-vision.md", "artifacts/pm/01-charter.md", "artifacts/ba/02-brd.md", "traceability/rtm.md"]
    return [
        "artifacts/stage-0/project-brief.md",
        "artifacts/po/01-vision.md",
        "artifacts/pm/01-charter.md",
        "artifacts/ba/01-prd.md",
        "artifacts/ba/02-brd.md",
        "artifacts/ba/03-srs/srs.md",
        "artifacts/ba/05-test-cases.md",
        "artifacts/ic/03-deployment-plan.md",
        "artifacts/ic/04-uat-plan.md",
        "traceability/rtm.md",
    ]


def _profile_note(profile: str) -> str:
    notes = {
        "client-ready": "External review pack: hides raw sources, prioritizes business context, scope, SRS, UAT and traceability.",
        "internal": "Internal working pack: includes PM/PO/BA/IC artifacts for delivery coordination.",
        "developer": "Developer handoff: prioritizes SRS, APIs, workflows, test cases and deployment notes.",
        "management": "Executive pack: focuses on vision, charter, BRD and traceability summary.",
    }
    return notes.get(profile, "Custom export profile.")

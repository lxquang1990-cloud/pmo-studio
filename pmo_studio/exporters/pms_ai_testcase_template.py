"""PMS AI Q&A testcase workbook template/exporter.

BA/QA-facing standard:
- 00_Bia
- 00_TomTat_KetQua_Test
- 01_ChuThich
- 02_TatCa_TestCases

Raw execution/debug data must live in a separate machine-facing artifact, not in
this workbook. The workbook is optimized for stakeholder review and manual QA.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Any
import json
import re

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CANONICAL_COLUMNS = [
    "Mã TC",
    "Module",
    "Chức năng",
    "Loại test",
    "Priority",
    "Tiêu đề testcase",
    "Tiền điều kiện",
    "Các bước thực hiện",
    "Dữ liệu test",
    "Kết quả mong đợi",
    "Kết quả thực tế",
    "Trạng thái",
    "Tester",
    "Ngày test",
    "Ghi chú",
]

REMOVED_TECHNICAL_COLUMNS = [
    "Mã defect",
    "HTTP Status",
    "API Success",
    "Actual response",
    "Follow-up questions",
    "Auto verdict",
    "Execution note",
    "Executed at",
]

GROUP_LABELS = {
    "02_Nhom_A_Factual": "A - Dữ kiện trực tiếp",
    "03_Nhom_B_TongHop": "B - Tổng hợp, tính toán",
    "04_Nhom_C_TimKiem": "C - Tìm kiếm, lọc dữ liệu",
    "05_Nhom_D_Timeline": "D - Thời gian, tiến trình",
    "06_Nhom_E_Nested": "E - Dữ liệu lồng nhau/chi tiết",
    "07_Nhom_F_File": "F - File đính kèm",
    "08_Nhom_G_SoSanh": "G - So sánh",
    "09_Nhom_H_Citation": "H - Citation/nguồn tham chiếu",
    "10_Nhom_I_PhanQuyen": "I - Phân quyền dữ liệu",
    "11_Nhom_J_Adversarial": "J - Adversarial/chống bịa dữ liệu",
    "FACT": "A - Dữ kiện trực tiếp",
    "AGG": "B - Tổng hợp, tính toán",
    "FIND": "C - Tìm kiếm, lọc dữ liệu",
    "TIME": "D - Thời gian, tiến trình",
    "NESTED": "E - Dữ liệu lồng nhau/chi tiết",
    "FILE": "F - File đính kèm",
    "CMP": "G - So sánh",
    "CITE": "H - Citation/nguồn tham chiếu",
    "PERM": "I - Phân quyền dữ liệu",
    "ADV": "J - Adversarial/chống bịa dữ liệu",
}

DEFAULT_GROUPS = [
    ("02_Nhom_A_Factual", "Nhóm A - Câu hỏi dữ kiện trực tiếp", "FACT"),
    ("03_Nhom_B_TongHop", "Nhóm B - Câu hỏi tổng hợp và tính toán", "AGG"),
    ("04_Nhom_C_TimKiem", "Nhóm C - Câu hỏi tìm kiếm và lọc dữ liệu", "FIND"),
    ("05_Nhom_D_Timeline", "Nhóm D - Câu hỏi theo thời gian và tiến trình", "TIME"),
    ("06_Nhom_E_Nested", "Nhóm E - Câu hỏi dữ liệu lồng nhau/chi tiết", "NESTED"),
    ("07_Nhom_F_File", "Nhóm F - Câu hỏi về file đính kèm", "FILE"),
    ("08_Nhom_G_SoSanh", "Nhóm G - Câu hỏi so sánh", "CMP"),
    ("09_Nhom_H_Citation", "Nhóm H - Kiểm tra citation/nguồn tham chiếu", "CITE"),
    ("10_Nhom_I_PhanQuyen", "Nhóm I - Kiểm tra phân quyền dữ liệu", "PERM"),
    ("11_Nhom_J_Adversarial", "Nhóm J - Kiểm tra adversarial/chống bịa dữ liệu", "ADV"),
]

TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates" / "excel" / "pms_ai_qa_testcase_template.xlsx"

@dataclass
class PmsAiTestCase:
    ma_tc: str
    module: str
    chuc_nang: str
    loai_test: str
    priority: str
    tieu_de: str
    tien_dieu_kien: str
    cac_buoc: str
    du_lieu_test: str
    ket_qua_mong_doi: str
    ket_qua_thuc_te: str = ""
    trang_thai: str = "Not Run"
    tester: str = ""
    ngay_test: str = ""
    ma_defect: str = ""  # kept for backwards compatibility; intentionally not exported
    ghi_chu: str = ""
    group_sheet: str = "02_Nhom_A_Factual"

    def normalized_loai_test(self) -> str:
        return GROUP_LABELS.get(self.group_sheet) or GROUP_LABELS.get(self.loai_test) or self.loai_test

    def to_row(self) -> list[str]:
        return [
            self.ma_tc,
            self.module,
            self.chuc_nang,
            self.normalized_loai_test(),
            self.priority,
            self.tieu_de,
            self.tien_dieu_kien,
            self.cac_buoc,
            self.du_lieu_test,
            self.ket_qua_mong_doi,
            extract_visible_ai_answer(self.ket_qua_thuc_te),
            normalize_status(self.trang_thai),
            self.tester,
            self.ngay_test,
            clean_internal_text(self.ghi_chu),
        ]

def normalize_status(value: Any) -> str:
    text = str(value or "Not Run").strip()
    low = text.lower()
    if low in {"pass", "passed", "pass_manual", "pass_auto"}:
        return "Pass"
    if low in {"fail", "failed", "fail_manual", "fail_auto", "fail_api"}:
        return "Fail"
    if low in {"review", "need review"}:
        return "Review"
    if low in {"blocked", "block"}:
        return "Blocked"
    if low in {"", "not run", "not_run", "ready"}:
        return "Not Run"
    return text

def extract_visible_ai_answer(value: Any) -> str:
    """Return only the user-visible AI answer, stripping API/log wrappers."""
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    if "AI answer:" in text:
        text = text.split("AI answer:", 1)[1].strip()
    if text.startswith("{"):
        try:
            obj = json.loads(text)
            answer = obj.get("data", {}).get("answer") or obj.get("answer")
            if answer:
                return str(answer).strip()
        except Exception:
            pass
    lines = []
    for line in text.splitlines():
        if line.strip().startswith(("Executed at:", "HTTP status:", "Elapsed:", "AI answer:")):
            continue
        lines.append(line)
    return "\n".join(lines).strip()

def clean_internal_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    patterns = [
        r"Generated từ danh sách câu hỏi end-user thống kê PMS AI\..*?(?:\n|$)",
        r"Đã gộp nhóm A-J vào một sheet\..*?J-Adversarial\.?(?:\n|$)",
        r"Imported from latest executed template.*?(?:\n|$)",
        r"HTTP/API raw fields.*?(?:\n|$)",
    ]
    for pattern in patterns:
        text = re.sub(pattern, "", text, flags=re.I | re.S)
    return re.sub(r"\n{3,}", "\n\n", text).strip()

def analyze_template(path: Path = TEMPLATE_PATH) -> dict:
    """Return the current BA/QA-facing PMS AI testcase schema summary."""
    workbook_path = path
    if path.exists():
        try:
            wb = load_workbook(path, data_only=False)
            sheets = []
            for ws in wb.worksheets:
                header_row = None
                headers: list[str] = []
                for row in range(1, min(ws.max_row, 10) + 1):
                    values = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
                    if values[: len(CANONICAL_COLUMNS)] == CANONICAL_COLUMNS:
                        header_row = row
                        headers = values[: len(CANONICAL_COLUMNS)]
                        break
                sheets.append({"name": ws.title, "rows": ws.max_row, "cols": ws.max_column, "header_row": header_row, "headers": headers})
            return {"template": str(workbook_path), "sheets": sheets, "canonical_columns": CANONICAL_COLUMNS, "removed_technical_columns": REMOVED_TECHNICAL_COLUMNS}
        except Exception:
            pass
    return {"template": str(workbook_path), "sheets": [], "canonical_columns": CANONICAL_COLUMNS, "removed_technical_columns": REMOVED_TECHNICAL_COLUMNS}

def _styles():
    thin = Side(style="thin", color="D9D9D9")
    return {
        "blue": "1F4E78",
        "mid_blue": "5B9BD5",
        "light_blue": "D9EAF7",
        "white": "FFFFFF",
        "green": "E2F0D9",
        "red": "FCE4D6",
        "yellow": "FFF2CC",
        "grey": "FAFAFA",
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
    }

def _write_key_value_sheet(ws, rows: list[tuple[Any, Any]]) -> None:
    s = _styles()
    ws.sheet_view.showGridLines = False
    for r, (a, b) in enumerate(rows, 1):
        ws.cell(r, 1, a)
        ws.cell(r, 2, b)
        for c in (1, 2):
            cell = ws.cell(r, c)
            cell.border = s["border"]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=s["light_blue"])
    for c in (1, 2):
        ws.cell(1, c).font = Font(bold=True, color=s["white"])
        ws.cell(1, c).fill = PatternFill("solid", fgColor=s["blue"])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    ws.freeze_panes = "A2"

def _style_main_sheet(ws, row_count: int) -> None:
    s = _styles()
    ws.sheet_view.showGridLines = False
    max_col = len(CANONICAL_COLUMNS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws.cell(1, 1, "TẤT CẢ TEST CASES PMS AI Q&A")
    ws.cell(1, 1).font = Font(bold=True, size=14, color=s["white"])
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=s["blue"])
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    ws.cell(2, 1, "Danh sách testcase đã gộp. Sử dụng bộ lọc ở hàng tiêu đề để kiểm tra theo Loại test, Priority hoặc Trạng thái.")
    ws.cell(2, 1).font = Font(italic=True, color="666666")
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    for c, header in enumerate(CANONICAL_COLUMNS, 1):
        cell = ws.cell(3, c, header)
        cell.font = Font(bold=True, color=s["white"])
        cell.fill = PatternFill("solid", fgColor=s["mid_blue"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = s["border"]
    status_col = CANONICAL_COLUMNS.index("Trạng thái") + 1
    priority_col = CANONICAL_COLUMNS.index("Priority") + 1
    for r in range(4, 4 + row_count):
        alt = PatternFill("solid", fgColor=s["grey"]) if (r - 3) % 2 == 0 else None
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            cell.border = s["border"]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if alt:
                cell.fill = alt
        st = ws.cell(r, status_col).value
        ws.cell(r, status_col).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, status_col).font = Font(bold=True)
        if st == "Pass":
            ws.cell(r, status_col).fill = PatternFill("solid", fgColor=s["green"])
        elif st == "Fail":
            ws.cell(r, status_col).fill = PatternFill("solid", fgColor=s["red"])
        elif st == "Review":
            ws.cell(r, status_col).fill = PatternFill("solid", fgColor=s["yellow"])
        ws.cell(r, priority_col).alignment = Alignment(horizontal="center", vertical="center")
    widths = {
        "Mã TC": 22,
        "Module": 20,
        "Chức năng": 32,
        "Loại test": 34,
        "Priority": 10,
        "Tiêu đề testcase": 42,
        "Tiền điều kiện": 42,
        "Các bước thực hiện": 44,
        "Dữ liệu test": 44,
        "Kết quả mong đợi": 58,
        "Kết quả thực tế": 72,
        "Trạng thái": 14,
        "Tester": 14,
        "Ngày test": 14,
        "Ghi chú": 34,
    }
    for c, header in enumerate(CANONICAL_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths[header]
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[3].height = 36
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(max_col)}{max(3, 3 + row_count)}"

def _status_counts(cases: list[PmsAiTestCase]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for case in cases:
        status = normalize_status(case.trang_thai)
        counts[status] = counts.get(status, 0) + 1
    return counts

def export_pms_ai_testcases(
    cases: Iterable[PmsAiTestCase],
    out: Path,
    project: str = "PMS (Procurement Management System)",
    module: str = "AI Q&A - Hỏi đáp về dữ liệu",
    created_date: str = "",
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    """Export cases using the final BA/QA-facing PMS AI testcase format.

    ``template_path`` is accepted for backwards compatibility but the exporter now
    generates the standard 4-sheet workbook directly to guarantee schema quality.
    """
    case_list = list(cases)
    wb = Workbook()
    cover = wb.active
    cover.title = "00_Bia"
    _write_key_value_sheet(cover, [
        ("Thông tin", "Giá trị"),
        ("Tên bộ testcase", "PMS AI Q&A Test Cases"),
        ("Dự án/Hệ thống", project),
        ("Module", module),
        ("Số testcase", len(case_list)),
        ("Ngày tạo", created_date),
        ("Phạm vi", "Kiểm tra phản hồi AI Q&A cho nghiệp vụ PMS/eOffice."),
    ])

    summary = wb.create_sheet("00_TomTat_KetQua_Test")
    counts = _status_counts(case_list)
    summary_rows: list[tuple[Any, Any]] = [
        ("Hạng mục", "Nội dung"),
        ("Mục đích", "Bộ testcase kiểm tra phản hồi AI Q&A cho nghiệp vụ PMS/eOffice."),
        ("Cách sử dụng", "Mở sheet 02_TatCa_TestCases, lọc theo Loại test/Priority/Trạng thái, sau đó đối chiếu Câu hỏi – Kết quả mong đợi – Kết quả thực tế."),
        ("Sheet chính", "02_TatCa_TestCases"),
        ("Số testcase", len(case_list)),
        ("Cột kết quả test", "Trạng thái"),
        ("Cột phản hồi AI", "Kết quả thực tế"),
        ("Ghi chú", "File BA/QA-facing; raw API/debug execution nên lưu ở artifact riêng."),
    ]
    for status in ["Pass", "Fail", "Review", "Not Run", "Blocked"]:
        if status in counts:
            summary_rows.append((f"Trạng thái {status}", counts[status]))
    _write_key_value_sheet(summary, summary_rows)

    legend = wb.create_sheet("01_ChuThich")
    _write_key_value_sheet(legend, [
        ("Nội dung", "Mô tả"),
        ("Loại test", "Nhóm kiểm tra nghiệp vụ, ví dụ: A - Dữ kiện trực tiếp, B - Tổng hợp/tính toán, I - Phân quyền dữ liệu."),
        ("Priority", "Mức ưu tiên kiểm tra. P1 là nhóm cần ưu tiên cao."),
        ("Kết quả mong đợi", "Tiêu chí để đánh giá phản hồi AI là đúng/chấp nhận được."),
        ("Kết quả thực tế", "Nội dung phản hồi thực tế của AI sau khi chạy test."),
        ("Trạng thái", "Kết quả kiểm tra: Pass / Fail / Review / Not Run / Blocked."),
        ("Tester / Ngày test", "Người thực hiện kiểm tra và ngày ghi nhận kết quả."),
        ("Ghi chú", "Thông tin bổ sung cho BA/QA nếu cần."),
    ])

    main = wb.create_sheet("02_TatCa_TestCases")
    for idx, case in enumerate(case_list, 4):
        for col, value in enumerate(case.to_row(), 1):
            main.cell(idx, col, value)
    _style_main_sheet(main, len(case_list))

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    _assert_final_workbook(out)
    return out

def _assert_final_workbook(path: Path) -> None:
    wb = load_workbook(path, read_only=False, data_only=True)
    expected_sheets = ["00_Bia", "00_TomTat_KetQua_Test", "01_ChuThich", "02_TatCa_TestCases"]
    assert wb.sheetnames == expected_sheets
    ws = wb["02_TatCa_TestCases"]
    headers = [ws.cell(3, c).value for c in range(1, len(CANONICAL_COLUMNS) + 1)]
    assert headers == CANONICAL_COLUMNS
    for bad in REMOVED_TECHNICAL_COLUMNS:
        assert bad not in headers
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if isinstance(cell.value, str):
                    assert "Generated từ danh sách câu hỏi" not in cell.value
                    assert "Đã gộp nhóm A-J vào một sheet" not in cell.value

def create_blank_pms_ai_template(out: Path, template_path: Path = TEMPLATE_PATH) -> Path:
    """Create a clean BA/QA-facing workbook with no testcase rows."""
    return export_pms_ai_testcases([], out, template_path=template_path)

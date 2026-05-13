"""Gate A structural checks for Markdown and Excel artifacts."""
from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook
from pmo_studio.core.ids import extract_ids, validate_id
from pmo_studio.rubrics.loader import load_rubric


@dataclass
class CheckResult:
    id: str
    passed: bool
    evidence: str
    blocker: bool = False


@dataclass
class GateResult:
    stage: str
    layer: str = "A"
    passed: bool = False
    checks: List[CheckResult] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {"stage": self.stage, "layer": self.layer, "passed": self.passed, "checks": [asdict(c) for c in self.checks]}


DEFAULT_MARKDOWN_SECTIONS = {
    "stage-0.project_brief": ["# Project Brief", "## 1. Mục tiêu kinh doanh", "## 3. Phạm vi cấp cao"],
    "ba.brd": ["# BRD", "## 1. Business Context", "## 2. Business Requirements"],
    "ba.srs": ["# SRS", "## 1. Introduction", "## 2. Overall Description", "## 3. Specific Requirements"],
    "ba.us_ac": ["# User Story", "## Acceptance Criteria"],
}

EXCEL_REQUIRED = {
    "ba.quotation": {
        "Feature List": ["STT", "Chức năng / Màn hình", "Đơn giá (Man/day)", "Manday", "Thành tiền (VNĐ)"],
        "Giả định":    ["STT", "Loại", "Nội dung giả định"],
    },
    "ic.fit_gap": {
        "Summary": ["Metric", "Value"],
        "Fit-Gap Detail": ["BR/REQ ID", "Description", "Fit/Gap", "Action", "Effort (md)"]
    },
    "ic.config_workbook": {
        "Summary": ["Metric", "Value"],
        "Org Settings": ["Parameter", "Value", "Required", "Owner", "Status"]
    }
}

EXCEL_ENUMS = {
    "ic.fit_gap": {"Fit-Gap Detail": {"Fit/Gap": {"Fit", "Gap"}, "Gap Type": {"Fit", "Customization", "Workaround", "Out of Scope", ""}}},
    "ic.config_workbook": {"Org Settings": {"Required": {"yes", "no", "yes (if SSO_AD)", ""}, "Status": {"Draft", "Pending", "Confirmed", ""}}},
    "ba.quotation": {"Giả định": {"Loại": {"Phạm vi", "Kỹ thuật", "Dữ liệu"}}},
}


def _heading_exists(text: str, heading: str) -> bool:
    return any(line.strip().startswith(heading) for line in text.splitlines())


def check_markdown(path: Path, stage: str) -> GateResult:
    text = path.read_text(encoding="utf-8")
    checks: list[CheckResult] = []
    rubric = load_rubric(stage, "A") or {}
    sections = rubric.get("sections") or DEFAULT_MARKDOWN_SECTIONS.get(stage, [])
    for heading in sections:
        checks.append(CheckResult(f"section.{heading.lower().replace(' ', '_')}", _heading_exists(text, heading), f"Heading {heading}"))
    ids = extract_ids(text)
    invalid = [i for i in re.findall(r"\b[A-Z]{2,8}(?:-[A-Z]{2,8})?-\d{3}(?:-\d{2})?\b", text) if not validate_id(i)]
    checks.append(CheckResult("id.valid", not invalid, "All IDs valid" if not invalid else f"Invalid IDs: {invalid[:10]}", True))
    dupes = sorted({i for i in ids if ids.count(i) > 1})
    checks.append(CheckResult("id.no_duplicates", not dupes, "No duplicate IDs" if not dupes else f"Duplicates: {dupes[:10]}"))
    result = GateResult(stage=stage, checks=checks)
    result.passed = _passed(checks)
    return result


def _header_values(ws) -> list[str]:
    return [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]


def _rows_as_dicts(ws) -> tuple[list[str], list[dict[str, Any]]]:
    headers = _header_values(ws)
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return headers, rows


def check_excel(path: Path, stage: str) -> GateResult:
    wb = load_workbook(path, data_only=False)
    checks: list[CheckResult] = []
    required = EXCEL_REQUIRED.get(stage, {})
    known_ids = _collect_workbook_ids(wb)
    for sheet, cols in required.items():
        exists = sheet in wb.sheetnames
        checks.append(CheckResult(f"sheet.{sheet}", exists, f"Sheet {sheet} {'exists' if exists else 'missing'}", True))
        if not exists:
            continue
        ws = wb[sheet]
        headers, rows = _rows_as_dicts(ws)
        missing = [c for c in cols if c not in headers]
        checks.append(CheckResult(f"columns.{sheet}", not missing, "Required columns present" if not missing else f"Missing: {missing}", True))
        if missing:
            continue
        empty = _empty_mandatory(rows, cols)
        checks.append(CheckResult(f"mandatory.{sheet}", not empty, "Mandatory cells filled" if not empty else f"Empty mandatory cells: {empty[:10]}", True))
        enum_errors = _enum_errors(stage, sheet, rows)
        checks.append(CheckResult(f"enum.{sheet}", not enum_errors, "Enum values valid" if not enum_errors else f"Invalid enum: {enum_errors[:10]}"))
        id_errors = _id_errors(rows)
        checks.append(CheckResult(f"ids.{sheet}", not id_errors, "Referenced IDs valid" if not id_errors else f"Invalid IDs: {id_errors[:10]}", True))
        if stage == "ba.quotation" and sheet == "Feature List":
            checks.extend(_quotation_checks_v4(wb, rows))
        if stage == "ic.config_workbook" and sheet == "Org Settings":
            plain = _plain_secret_errors(rows)
            checks.append(CheckResult("secret.org_settings", not plain, "No plaintext secrets" if not plain else f"Plain secret-like values: {plain[:10]}", True))
    result = GateResult(stage=stage, checks=checks)
    result.passed = _passed(checks)
    return result


def _collect_workbook_ids(wb) -> set[str]:
    ids: set[str] = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if value is not None:
                    ids.update(extract_ids(str(value)))
    return ids


def _empty_mandatory(rows: list[dict[str, Any]], cols: list[str]) -> list[str]:
    errors = []
    for idx, row in enumerate(rows, start=2):
        if not any(v not in {None, ""} for v in row.values()):
            continue
        for col in cols:
            if row.get(col) in {None, ""}:
                errors.append(f"row {idx} col {col}")
    return errors


def _enum_errors(stage: str, sheet: str, rows: list[dict[str, Any]]) -> list[str]:
    errors = []
    enums = EXCEL_ENUMS.get(stage, {}).get(sheet, {})
    for idx, row in enumerate(rows, start=2):
        for col, allowed in enums.items():
            if col in row:
                val = "" if row.get(col) is None else str(row.get(col)).strip()
                if val not in allowed:
                    errors.append(f"row {idx} {col}={val}")
    return errors


def _id_errors(rows: list[dict[str, Any]]) -> list[str]:
    errors = []
    for idx, row in enumerate(rows, start=2):
        for col, val in row.items():
            if val is None:
                continue
            for token in re.findall(r"\b[A-Z]{2,8}(?:-[A-Z]{2,8})?-\d{3}(?:-\d{2})?\b", str(val)):
                if not validate_id(token):
                    errors.append(f"row {idx} {col}={token}")
    return errors


def _quotation_checks(wb, rows: list[dict[str, Any]]) -> list[CheckResult]:
    checks = []
    detail_total_md = 0.0
    detail_total_cost = 0.0
    numeric_errors = []
    for idx, row in enumerate(rows, start=2):
        try:
            md = float(row.get("Total md") or 0)
            cost = float(row.get("Cost VND") or 0)
            detail_total_md += md
            detail_total_cost += cost
        except Exception:
            numeric_errors.append(f"row {idx}")
    checks.append(CheckResult("quotation.numeric", not numeric_errors, "Numeric totals valid" if not numeric_errors else f"Non-numeric totals: {numeric_errors[:10]}", True))
    summary = _summary_metrics(wb["Summary"]) if "Summary" in wb.sheetnames else {}
    if "Total manday" in summary:
        checks.append(CheckResult("quotation.total_manday", abs(float(summary["Total manday"]) - detail_total_md) < 0.01, f"Summary={summary['Total manday']} Detail={detail_total_md}", True))
    if "Total cost VND" in summary:
        checks.append(CheckResult("quotation.total_cost", abs(float(summary["Total cost VND"]) - detail_total_cost) < 1, f"Summary={summary['Total cost VND']} Detail={detail_total_cost}", True))
    return checks


def _quotation_checks_v4(wb, rows: list[dict[str, Any]]) -> list[CheckResult]:
    """Gate checks cho BaoGia_Template_v4 (Feature List sheet)."""
    checks = []
    manday_col = "Manday"
    thanhtien_col = "Thành tiền (VNĐ)"
    numeric_errors = []
    has_screens = False
    for idx, row in enumerate(rows, start=2):
        md_val = row.get(manday_col)
        tt_val = row.get(thanhtien_col)
        if md_val is not None and str(md_val).startswith("="):
            continue  # formula — skip validation (không data_only)
        try:
            if md_val is not None and str(md_val).strip() not in ("", "Manday"):
                float(md_val)
                has_screens = True
            if tt_val is not None and str(tt_val).strip() not in ("", "Thành tiền (VNĐ)"):
                float(str(tt_val).replace(",", ""))
        except Exception:
            numeric_errors.append(f"row {idx}")
    checks.append(CheckResult("quotation.v4.numeric", not numeric_errors,
                              "Numeric cells valid" if not numeric_errors else f"Non-numeric: {numeric_errors[:10]}",
                              blocker=False))
    checks.append(CheckResult("quotation.v4.has_screens", has_screens,
                              "Feature List has screen data", blocker=True))
    has_tonghop = "Tổng hợp" in wb.sheetnames
    has_giadinh = "Giả định" in wb.sheetnames
    checks.append(CheckResult("quotation.v4.sheet_tonghop", has_tonghop,
                              "Sheet Tổng hợp exists" if has_tonghop else "Sheet Tổng hợp missing",
                              blocker=True))
    checks.append(CheckResult("quotation.v4.sheet_giadinh", has_giadinh,
                              "Sheet Giả định exists" if has_giadinh else "Sheet Giả định missing",
                              blocker=True))
    return checks


def _summary_metrics(ws) -> dict[str, Any]:
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            out[str(row[0])] = row[1] if len(row) > 1 else None
    return out


def _plain_secret_errors(rows: list[dict[str, Any]]) -> list[str]:
    errors = []
    secret_re = re.compile(r"(?i)(password|secret|token|api[_-]?key)")
    for idx, row in enumerate(rows, start=2):
        param = str(row.get("Parameter") or "")
        value = str(row.get("Value") or "")
        if secret_re.search(param) and value and not value.startswith("vault:"):
            errors.append(f"row {idx} {param}")
    return errors


def _passed(checks: list[CheckResult]) -> bool:
    # Gate A is a structural safety gate: blocker checks fail the gate, while
    # non-blocker findings such as duplicate references remain visible evidence
    # for reviewers without stopping smoke generation.
    return all(c.passed for c in checks if c.blocker)


def run_gate_a(path: Path, stage: str) -> GateResult:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return check_excel(path, stage)
    return check_markdown(path, stage)


def save_gate_result(result: GateResult, project_root: Path) -> Path:
    out = project_root / "quality" / "gate-a" / f"{result.stage.replace('.', '-')}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result.to_dict(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return out

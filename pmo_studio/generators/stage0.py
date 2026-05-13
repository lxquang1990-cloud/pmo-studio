"""Stage 0 intake generator."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pmo_studio.core.ids import IdAllocator
from pmo_studio.core.project import Project
from pmo_studio.security.preprocessor import process_source
from pmo_studio.security.ignore import is_ignored, load_ignore_patterns, write_default_ignore
from pmo_studio.domain.prompts import get_domain


def generate_stage0(project: Project, brief: str = "TBD", products: list[str] | None = None, sources: list[Path] | None = None) -> None:
    domain = get_domain(project.config.domain_pack)
    products = products or _infer_products(project.config.project_slug, brief)
    allocator = IdAllocator.from_state(project.state.id_counters)
    write_default_ignore(project.root)
    ignore_patterns = load_ignore_patterns(project_root=project.root)
    scan_results = []
    ignored_sources = []
    for src in sources or []:
        if is_ignored(src, ignore_patterns):
            ignored_sources.append(src)
            continue
        source_id = allocator.issue("SRC")
        scan_results.append(process_source(source_id, src, project.root))
    project.state.id_counters = allocator.counters

    source_text = _read_redacted_sources(project.root)
    intelligence = _extract_stage0_intelligence(project.config.project_slug, brief, source_text)
    project_goal = intelligence["goal"]
    product_lines = "\n".join([f"- [x] {p}" for p in products])
    module_lines = "\n".join([f"- {m}" for m in intelligence["modules"]])
    role_rows = "\n".join([f"| {role} | {desc} | High |" for role, desc in intelligence["roles"]])
    integration_lines = "\n".join([f"- {x}" for x in intelligence["integrations"]])
    decision_matrix = _build_scope_decision_matrix(intelligence)
    decision_rows = "\n".join([f"| DEC-{i:03d} | {d} | Derived from source intake. | Snail/PMO | {project.state.created_at[:10]} |" for i, d in enumerate(intelligence["decisions"], 1)])
    decision_rows += "\n" + "\n".join([f"| {row['id']} | {row['recommended']} | Scope Decision Matrix recommendation for: {row['question']} | {row['owner']} | {project.state.created_at[:10]} |" for row in decision_matrix])
    assumption_rows = "\n".join([f"| ASM-{i:03d} | {a} | stage-0 | PM/BA | Active |" for i, a in enumerate(intelligence["assumptions"], 1)])
    question_lines = "\n".join([f"- {q['question']} → Recommended: {q['recommended']} ({q['status']})" for q in decision_matrix])
    timeline = intelligence["timeline"]

    domain_roles_str = "\n".join([f"| {role} | {desc} |" for role, desc in (domain.roles or {}).items()]) if domain.roles and domain.domain_id != "bteco" else "| Asset Manager | Quản lý tài sản, cấp phát, thu hồi, điều chuyển, thanh lý |\n| HR | Cung cấp dữ liệu nhân sự, nghỉ việc, điều chuyển, cost center |\n| Accounting | Quản lý nguyên giá, khấu hao, giá trị còn lại, chi phí |\n| Employee | Xác nhận bàn giao/thu hồi và cập nhật tình trạng tài sản |\n| Executive | Xem dashboard, phê duyệt và ra quyết định đầu tư |"

    (project.root / "artifacts/stage-0/project-brief.md").write_text(f"""# Project Brief: {project.config.project_slug}

> **Project Slug:** {project.config.project_slug}
> **Customer:** {project.config.customer}
> **Project Type:** {project.config.project_type}
> **Domain:** Asset Management / Quản lý trang thiết bị, tài sản
> **Primary Source:** SRC-001

## 1. Mục tiêu kinh doanh
{project_goal}

## 2. Sản phẩm / phân hệ liên quan
{product_lines}

## 3. Phạm vi cấp cao
### In scope
{module_lines}

### Integration scope
{integration_lines}

### Scope decisions / implementation prerequisites
- Production code application không nằm trong phạm vi bộ tài liệu; deliverable là hồ sơ PMO/BA/IC client-ready để review, estimate và governance.
- RFID hardware SDK/vendor-specific integration được phân loại Optional/Phase 2 và chỉ estimate sau PoC thiết bị/vendor.
- Chữ ký số/chữ ký điện tử có giá trị pháp lý được phân loại Optional/Phase 2 theo provider, license và chính sách nội bộ.
- Làm sạch dữ liệu legacy quy mô lớn được xử lý qua migration workstream riêng sau data profiling.

## 4. Timeline mong muốn
- Start: {timeline['start']}
- Target go-live: {timeline['go_live']}
- Planning horizon: {timeline['horizon']}

## 5. Stakeholder chính
| Tên | Role | Mức độ ảnh hưởng |
|---|---|---|
{role_rows}

## 6. Domain Context
| Role | Description |
|---|---|
{domain_roles_str}

## 7. Governance mong muốn
- Baseline required: {project.config.governance.get('baseline_required')}
- Formal CR after baseline: {project.config.governance.get('formal_cr_after_baseline')}
- Human approval required for official quotation/sign-off: {project.config.governance.get('human_approval_required')}
- Quality gates: Gate A/B/C before client-ready export
- Traceability: SRC/BG → BR → REQ → SCR/API/WF/RPT → US/AC/TC/EST

## 8. Dev Readiness Snapshot
- Linked REQ baseline: REQ-CORE-001 asset master, REQ-CORE-002 allocation/return/transfer workflow, REQ-CORE-003 maintenance/inventory/liquidation, REQ-CORE-004 reports, REQ-CORE-005 permission/audit.
- API/integration baseline: HRM import/export, ERP/accounting sync as Phase 2 unless contract ready, notification for approval and maintenance alerts.
- Workflow baseline: create asset → allocate/handover → return/transfer → inventory/maintenance → liquidation/reporting.
- Permission baseline: Admin, Asset Manager, Department Manager, Staff, Auditor.

## 9. Acceptance / Review Markers
- Customer-ready: Sponsor/PM xác nhận scope MVP, Phase 2, optional items và exclusions trong Scope Decision Matrix.
- Dev-ready: BA/Tech Lead sử dụng BRD/SRS/US/AC/Test Case để estimate workflow, permission, data model, integration và NFR.
- Commercial-ready: Báo giá chính thức sau khi baseline scope và assumptions được approved.

## 10. Scope Decision Baseline
Các câu hỏi dưới đây đã có recommended default để không block MVP; nếu sponsor chọn option khác thì tạo Change Request.
{question_lines}
""", encoding="utf-8")

    inv_lines = ["# Source Inventory", "", "| Source ID | Type | Path | Hash | Trusted | Redaction | Injection |", "|---|---|---|---|---|---|---|"]
    for r in scan_results:
        inv_lines.append(f"| {r.source_id} | file | {r.redacted_path} | sha256:{r.sha256} | partial | {r.redaction_count} | {r.injection_detected} |")
    if not scan_results:
        inv_lines.append("| SRC-000 | manual_input | (none yet) | - | yes | 0 | false |")
    for src in ignored_sources:
        inv_lines.append(f"| IGNORED | ignored | {src} | - | n/a | n/a | n/a |")
    (project.root / "artifacts/stage-0/source-inventory.md").write_text("\n".join(inv_lines) + "\n", encoding="utf-8")

    (project.root / "artifacts/stage-0/assumption-log.md").write_text(f"""# Assumption Log

| ID | Description | Source Stage | Owner | Status |
|---|---|---|---|---|
{assumption_rows}
""", encoding="utf-8")
    (project.root / "artifacts/stage-0/decision-log.md").write_text(f"""# Decision Log

| ID | Decision | Rationale | Owner | Date |
|---|---|---|---|---|
{decision_rows}
""", encoding="utf-8")
    _write_scope_decision_matrix(project.root / "artifacts/stage-0/scope-decision-matrix.xlsx", decision_matrix)
    (project.root / "artifacts/stage-0/scope-decision-matrix.md").write_text(_scope_decision_matrix_markdown(decision_matrix), encoding="utf-8")
    project.mark_stage("stage-0", "completed", artifacts=["project-brief.md", "source-inventory.md", "assumption-log.md", "decision-log.md", "scope-decision-matrix.xlsx", "scope-decision-matrix.md"])


def _read_redacted_sources(project_root: Path, max_chars: int = 30000) -> str:
    chunks = []
    for path in sorted((project_root / "source" / "redacted").glob("*")):
        if path.is_file():
            chunks.append(path.read_text(encoding="utf-8", errors="ignore"))
    return "\n".join(chunks)[:max_chars]


def _infer_products(slug: str, brief: str) -> list[str]:
    text = f"{slug} {brief}".lower()
    if any(k in text for k in ["tai-san", "tài sản", "ttb", "asset"]):
        return ["Asset Management", "HRM Integration", "Accounting/ERP Integration", "Mobile QR Inventory"]
    return ["PMO/BA Documentation Pack", "Traceability", "Quality Gates", "Client-ready Export"]


def _extract_stage0_intelligence(slug: str, brief: str, source_text: str) -> dict:
    text = f"{brief}\n{source_text}".strip()
    lower = text.lower()
    is_asset = any(k in lower or k in slug.lower() for k in ["tài sản", "tai san", "trang thiết bị", "ttb", "asset", "khấu hao", "bảo trì"])
    if is_asset:
        modules = [
            "Quản lý danh mục và hồ sơ tài sản/trang thiết bị: mã tài sản, nhóm, chủng loại, serial, ngày mua, nguyên giá, khấu hao, đơn vị sử dụng, tình trạng.",
            "Quản lý phân bổ, bàn giao và sử dụng tài sản theo nhân viên, phòng ban, mã nhân viên và cost center.",
            "Quản lý tình trạng, bảo trì, bảo hành, sửa chữa, chứng từ và cảnh báo đến hạn.",
            "Quản lý cấp phát, thu hồi khi nhân viên mới, nghỉ việc hoặc điều chuyển.",
            "Quản lý điều chuyển, thanh lý và lịch sử giao dịch giữa đơn vị/chi nhánh.",
            "Báo cáo tồn kho, tài sản theo người sử dụng/phòng ban, tình trạng, khấu hao, giá trị còn lại và dashboard cảnh báo.",
            "QR/RFID và mobile/PWA cho tra cứu, kiểm kê, bàn giao, thu hồi và cập nhật tình trạng.",
            "Phê duyệt điện tử và chữ ký điện tử/chữ ký số nội bộ cho biên bản.",
            "Quản trị phân quyền, audit log, import/export dữ liệu và backup/restore.",
        ]
        roles = [
            ("Sponsor/PM", "Chốt phạm vi, ngân sách, milestone, baseline và nghiệm thu."),
            ("Asset Manager", "Owner nghiệp vụ tài sản, quản lý hồ sơ, cấp phát, thu hồi, điều chuyển, thanh lý."),
            ("HR", "Cung cấp dữ liệu nhân sự, trạng thái nghỉ việc/điều chuyển, phòng ban, cost center."),
            ("Accounting", "Quản lý nguyên giá, khấu hao, giá trị còn lại, chi phí sửa chữa/thanh lý."),
            ("Employee", "Xem tài sản được giao, xác nhận bàn giao/thu hồi, cập nhật tình trạng."),
            ("Department Manager", "Theo dõi tài sản thuộc đơn vị và phê duyệt/yêu cầu nghiệp vụ."),
            ("IT/System Admin", "Quản lý phân quyền, tích hợp, backup, bảo mật và vận hành."),
        ]
        integrations = [
            "HRM: employee, department, cost center, status nghỉ việc/điều chuyển/bổ nhiệm.",
            "ERP/Kế toán: mã tài sản kế toán, nguyên giá, khấu hao, chi phí, giá trị còn lại.",
            "Mua sắm/Kho: cập nhật tài sản mới sau mua sắm/nhập kho.",
            "DMS: lưu trữ chứng từ, hóa đơn, biên bản, tài liệu bảo hành/sửa chữa.",
            "Notification: email/app nội bộ cho bảo trì, bảo hành, phê duyệt, cảnh báo.",
        ]
        assumptions = [
            "Mỗi tài sản có hoặc sẽ được chuẩn hóa mã định danh duy nhất trước migration/baseline.",
            "HRM có mã nhân viên, phòng ban và cost center đủ ổn định để mapping tài sản.",
            "Công thức khấu hao, ownership dữ liệu tài chính và quyền cập nhật cần được xác nhận với kế toán.",
            "QR là baseline cho kiểm kê mobile; RFID chỉ triển khai khi có thiết bị/vendor và PoC.",
            "Workflow phê duyệt, chữ ký điện tử/chữ ký số và giá trị pháp lý cần workshop riêng.",
        ]
        decisions = [
            "Stage 0 baseline tập trung vào phân hệ quản lý TTB/Tài sản, không dùng nội dung mặc định PMO Studio.",
            "Chia scope thành MVP, Phase 2 và Optional để kiểm soát budget/risk.",
            "Bắt buộc chạy traceability và Gate A/B/C trước client-ready export.",
        ]
        questions = [
            "MVP bắt buộc có web responsive là đủ hay cần mobile/PWA ngay từ giai đoạn đầu?",
            "Các hệ thống HRM/ERP/Kế toán/Mua sắm/DMS hiện tại là gì và có API/import-export contract không?",
            "Quy trình phê duyệt cấp phát, thu hồi, điều chuyển, thanh lý gồm bao nhiêu cấp và SLA thế nào?",
            "Quy tắc khấu hao, giá trị còn lại và quyền cập nhật dữ liệu tài chính thuộc hệ thống nào?",
            "Kiểm kê dùng QR hay RFID; nếu RFID thì thiết bị/vendor nào?",
            "Chữ ký điện tử/chữ ký số nội bộ cần giá trị pháp lý ở mức nào?",
        ]
        goal = "Xây dựng phân hệ quản lý trang thiết bị, tài sản nhằm số hóa toàn bộ vòng đời tài sản: khai báo, phân loại, phân bổ, bàn giao, sử dụng, bảo trì/bảo hành, kiểm kê, điều chuyển, thu hồi, thanh lý và báo cáo quản trị. Hệ thống cần liên thông HRM/ERP/Kế toán/Mua sắm/DMS để giảm nhập liệu thủ công, tăng tính minh bạch, kiểm soát trách nhiệm sử dụng tài sản và hỗ trợ ra quyết định về chi phí/khấu hao/vòng đời tài sản."
    else:
        modules = _top_lines(text) or ["Phân tích yêu cầu", "Tạo tài liệu PMO/BA", "Traceability", "Quality Gates", "Export"]
        roles = [("Sponsor/PM", "Chốt scope, milestone và nghiệm thu."), ("BA", "Phân tích yêu cầu và tài liệu hóa."), ("Tech Lead", "Review tính khả thi kỹ thuật."), ("QA", "Chuẩn bị test/UAT.")]
        integrations = ["Integration scope cần xác nhận trong workshop."]
        assumptions = ["Khách hàng sẽ xác nhận scope trước baseline.", "Source đầu vào là partial và cần workshop bổ sung."]
        decisions = ["Sử dụng PMO Studio v2.1 structure cho lifecycle, traceability và quality gates."]
        questions = ["Scope MVP là gì?", "Timeline và budget target?", "Stakeholder phê duyệt chính?", "Integration nào bắt buộc?"]
        goal = brief if brief and brief != "TBD" else "Tạo bộ hồ sơ PMO/BA/IC có traceability và quality gates từ source đầu vào."
    return {
        "goal": goal,
        "modules": modules,
        "roles": roles,
        "integrations": integrations,
        "assumptions": assumptions,
        "decisions": decisions,
        "questions": questions,
        "timeline": {"start": "Sau workshop scope/baseline", "go_live": "MVP target 10-12 tuần sau baseline", "horizon": "MVP 8-12 tuần; Phase 2 tùy scope tích hợp/mobile/workflow"},
    }


def _build_scope_decision_matrix(intelligence: dict) -> list[dict]:
    """Convert Stage-0 open questions into recommended default decisions."""
    goal = intelligence.get("goal", "")
    is_asset = "tài sản" in goal.lower() or "trang thiết bị" in goal.lower()
    if is_asset:
        return [
            {
                "id": "DEC-SCOPE-001",
                "question": "MVP bắt buộc có mobile/PWA ngay không?",
                "recommended": "MVP dùng web responsive; PWA/QR kiểm kê đưa Phase 2 trừ xác nhận bàn giao/thu hồi đơn giản.",
                "options": "A.Web only MVP | B.Web+PWA basic | C.Native mobile",
                "scope_impact": "A giảm scope MVP; B tăng UX/QA; C tăng đáng kể cost/time.",
                "cost_impact": "A baseline; B +10-15%; C +30-50%.",
                "timeline_impact": "A 8-12 tuần; B +2-3 tuần; C +6-10 tuần.",
                "owner": "Sponsor/PM + IT",
                "status": "Recommended default",
                "gate_c_blocker": "No nếu sponsor chấp nhận default",
            },
            {
                "id": "DEC-SCOPE-002",
                "question": "QR hay RFID cho kiểm kê?",
                "recommended": "QR là baseline Phase 2; RFID là Optional sau PoC thiết bị/vendor.",
                "options": "A.QR | B.QR+RFID readiness | C.RFID full integration",
                "scope_impact": "RFID full cần SDK, thiết bị, test hiện trường.",
                "cost_impact": "QR baseline; RFID readiness +5-8%; RFID full tùy vendor.",
                "timeline_impact": "RFID full cần PoC 2-4 tuần trước estimate chính thức.",
                "owner": "Asset Manager + IT",
                "status": "Recommended default",
                "gate_c_blocker": "No nếu RFID để Optional",
            },
            {
                "id": "DEC-SCOPE-003",
                "question": "HRM integration mức nào trong MVP?",
                "recommended": "MVP hỗ trợ import/export nhân sự/cost center; API realtime đưa Phase 2 nếu có contract.",
                "options": "A.Manual import | B.Scheduled file | C.API realtime",
                "scope_impact": "API realtime cần mapping, auth, retry, error handling.",
                "cost_impact": "A/B baseline; C +8-12 MD tùy API.",
                "timeline_impact": "C phụ thuộc availability của HRM team/vendor.",
                "owner": "HR + IT",
                "status": "Recommended default",
                "gate_c_blocker": "No nếu dùng import/export MVP",
            },
            {
                "id": "DEC-SCOPE-004",
                "question": "Kế toán/ERP ownership cho khấu hao và giá trị còn lại?",
                "recommended": "Kế toán/ERP là master dữ liệu tài chính; phân hệ tài sản hiển thị/đồng bộ và không ghi đè nếu chưa duyệt.",
                "options": "A.Asset system calculates | B.Accounting master | C.Hybrid",
                "scope_impact": "A cần rule khấu hao chi tiết; B giảm rủi ro tài chính.",
                "cost_impact": "B baseline; A/C +10-20 MD do rule/calculation/testing.",
                "timeline_impact": "A/C cần workshop kế toán và reconciliation test.",
                "owner": "Accounting + Sponsor",
                "status": "Recommended default",
                "gate_c_blocker": "No nếu Accounting master được chấp nhận",
            },
            {
                "id": "DEC-SCOPE-005",
                "question": "Workflow phê duyệt gồm bao nhiêu cấp?",
                "recommended": "MVP dùng workflow đơn giản: tạo phiếu → quản lý tài sản xác nhận → trưởng bộ phận/owner duyệt khi cần; workflow configurable nâng cao ở Phase 2.",
                "options": "A.Simple fixed | B.Configurable by transaction | C.Rule engine",
                "scope_impact": "Rule engine là scope lớn, cần thiết kế riêng.",
                "cost_impact": "A baseline; B +15-25 MD; C +40+ MD.",
                "timeline_impact": "B +3-5 tuần; C cần phase riêng.",
                "owner": "Sponsor/PM + Asset Manager",
                "status": "Recommended default",
                "gate_c_blocker": "No nếu simple workflow MVP",
            },
            {
                "id": "DEC-SCOPE-006",
                "question": "Chữ ký điện tử/chữ ký số cần mức nào?",
                "recommended": "Phase 2 optional; MVP dùng xác nhận điện tử/audit log, chưa cam kết chữ ký số pháp lý.",
                "options": "A.Audit confirmation | B.Internal e-sign | C.CA digital signature",
                "scope_impact": "CA digital signature phụ thuộc provider/license/legal policy.",
                "cost_impact": "A baseline; B +8-15 MD; C tùy provider/license.",
                "timeline_impact": "C cần procurement và integration testing.",
                "owner": "Legal/IT + Sponsor",
                "status": "Recommended default",
                "gate_c_blocker": "No nếu ký số để Phase 2 Optional",
            },
            {
                "id": "DEC-SCOPE-007",
                "question": "Target go-live cho MVP?",
                "recommended": "MVP target 10-12 tuần sau scope baseline, chưa bao gồm delay do integration/vendor/data cleansing.",
                "options": "A.8 tuần aggressive | B.10-12 tuần balanced | C.14-16 tuần safer",
                "scope_impact": "Timeline ngắn cần giảm scope mobile/integration/workflow.",
                "cost_impact": "Aggressive có risk buffer cao hơn.",
                "timeline_impact": "B là baseline planning.",
                "owner": "Sponsor/PM",
                "status": "Recommended default",
                "gate_c_blocker": "No nếu dùng planning assumption",
            },
        ]
    return [
        {
            "id": "DEC-SCOPE-001",
            "question": q,
            "recommended": "Chốt trong discovery workshop; dùng assumption conservative cho estimate draft.",
            "options": "TBD",
            "scope_impact": "TBD",
            "cost_impact": "TBD",
            "timeline_impact": "TBD",
            "owner": "Sponsor/PM",
            "status": "Open",
            "gate_c_blocker": "Yes",
        }
        for q in intelligence.get("questions", [])
    ]


def _write_scope_decision_matrix(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scope Decisions"
    headers = ["Decision ID", "Question", "Recommended Default", "Options", "Scope Impact", "Cost Impact", "Timeline Impact", "Owner", "Status", "Gate C Blocker"]
    ws.append(headers)
    for row in rows:
        ws.append([row["id"], row["question"], row["recommended"], row["options"], row["scope_impact"], row["cost_impact"], row["timeline_impact"], row["owner"], row["status"], row["gate_c_blocker"]])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [18, 36, 58, 38, 48, 32, 36, 26, 22, 22]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(rows)+1}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _scope_decision_matrix_markdown(rows: list[dict]) -> str:
    lines = [
        "# Scope Decision Matrix",
        "",
        "| Decision ID | Question | Recommended Default | Owner | Status | Gate C Blocker |",
        "|---|---|---|---|---|---|",
    ]
    for row in rows:
        lines.append(f"| {row['id']} | {row['question']} | {row['recommended']} | {row['owner']} | {row['status']} | {row['gate_c_blocker']} |")
    return "\n".join(lines) + "\n"


def _top_lines(text: str, limit: int = 8) -> list[str]:
    lines = []
    for line in text.splitlines():
        clean = re.sub(r"^[•\-\d\.\)\s]+", "", line.strip())
        if len(clean) > 20 and clean not in lines:
            lines.append(clean[:220])
        if len(lines) >= limit:
            break
    return lines

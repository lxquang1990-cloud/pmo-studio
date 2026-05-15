from __future__ import annotations
import hashlib, json, shutil, zipfile
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from importlib.metadata import version, PackageNotFoundError
from pmo_studio.model.generators import export_customer_from_model, ensure_model
from pmo_studio.model.quality import write_quality_v3

@dataclass
class DeliveryArtifact:
    kind: str
    source_path: str
    delivery_path: str
    sha256: str
    size_bytes: int

@dataclass
class DeliveryManifest:
    schema: str
    generated_at: str
    project_slug: str
    customer: str
    domain_id: str
    domain_label: str
    language: str
    pmo_studio_version: str
    model_hash: str
    quality_report: str
    artifacts: list[DeliveryArtifact]
    zip_path: str | None = None


def export_delivery_pack(project, lang: str = 'vi', *, make_zip: bool = True) -> Path:
    model = ensure_model(project)
    outputs = export_customer_from_model(project, lang)
    quality_json = write_quality_v3(project.root, scope='customer')
    model_path = project.root/'artifacts/model/ba-model.json'
    model_hash = _sha256(model_path)
    version = _version()
    delivery_root = project.root/f'exports/delivery/{lang}'
    if delivery_root.exists(): shutil.rmtree(delivery_root)
    delivery_root.mkdir(parents=True, exist_ok=True)
    artifacts=[]
    mapping={
        'srs_docx': project.root/f'exports/customer/{lang}/srs-customer-ready.docx',
        'user_stories_docx': project.root/f'exports/customer/{lang}/user-stories.docx',
        'uat_xlsx': project.root/f'exports/customer/{lang}/uat-pack/test-cases-uat.{lang}.xlsx',
        'quotation_xlsx': _select_customer_quotation(project, lang),
        'quality_json': quality_json,
        'quality_md': project.root/f'quality/customer-review-v3.customer.md',
        'ba_model_json': model_path,
    }
    for kind, src in mapping.items():
        if not src.exists(): continue
        dest = delivery_root / _delivery_name(kind, src, lang)
        shutil.copy2(src, dest)
        artifacts.append(DeliveryArtifact(kind, str(src), str(dest.relative_to(delivery_root)), _sha256(dest), dest.stat().st_size))
    manifest = DeliveryManifest('pmo.delivery_manifest.v1', datetime.now(timezone.utc).isoformat(), project.config.project_slug, project.config.customer, model.domain.id, model.domain.label, lang, version, model_hash, str(Path('customer-review-v3.customer.json')), artifacts)
    manifest_path = delivery_root/'DELIVERY_MANIFEST.json'
    manifest_path.write_text(json.dumps(asdict(manifest), ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    _write_manifest_md(delivery_root/'DELIVERY_MANIFEST.md', manifest)
    if make_zip:
        zip_path = delivery_root.parent/f'{project.config.project_slug}-customer-delivery-{lang}-v{version}.zip'
        if zip_path.exists(): zip_path.unlink()
        with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
            for p in delivery_root.rglob('*'):
                if p.is_file(): zf.write(p, arcname=f'{delivery_root.name}/{p.relative_to(delivery_root)}')
        manifest.zip_path = str(zip_path)
        manifest_path.write_text(json.dumps(asdict(manifest), ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
        _write_manifest_md(delivery_root/'DELIVERY_MANIFEST.md', manifest)
    return manifest_path



def _select_customer_quotation(project, lang: str) -> Path:
    """Pick the most customer-ready quotation workbook available.

    The model quote is intentionally compact. If a detailed customer quotation
    has been produced by a richer estimator/reviewer, delivery packs should use
    that detailed workbook instead of overwriting it with the compact model one.
    """
    candidates = [
        project.root/f'exports/customer/{lang}/quotation-customer-ready-detailed.{lang}.xlsx',
        project.root/f'exports/customer/{lang}/quotation-customer-ready-detailed.xlsx',
        project.root/f'exports/customer/{lang}/quotation-customer-ready.{lang}.xlsx',
        project.root/f'exports/customer/{lang}/quotation-customer-ready.xlsx',
    ]
    existing = [p for p in candidates if p.exists()]
    if not existing:
        return project.root/f'exports/customer/{lang}/quotation-customer-ready.{lang}.xlsx'
    return max(existing, key=_quotation_detail_score)

def _quotation_detail_score(path: Path) -> tuple[int, int]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheet_bonus = 50 if {'Feature List', 'Tổng hợp', 'Giả định'} <= set(wb.sheetnames) else 0
        rows = wb['Feature List'].max_row if 'Feature List' in wb.sheetnames else max(ws.max_row for ws in wb.worksheets)
        return (sheet_bonus + rows, path.stat().st_size)
    except Exception:
        return (0, path.stat().st_size)

def _delivery_name(kind: str, src: Path, lang: str) -> str:
    names={
        'srs_docx': f'01-srs-customer-ready.{lang}.docx',
        'user_stories_docx': f'02-user-stories.{lang}.docx',
        'uat_xlsx': f'03-test-cases-uat.{lang}.xlsx',
        'quotation_xlsx': f'04-quotation-customer-ready.{lang}.xlsx',
        'quality_json': 'customer-review-v3.customer.json',
        'quality_md': 'customer-review-v3.customer.md',
        'ba_model_json': 'ba-model.json',
    }
    return names.get(kind, src.name)

def _sha256(path: Path) -> str:
    h=hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024), b''): h.update(chunk)
    return h.hexdigest()

def _version() -> str:
    import tomllib
    pyproject=Path(__file__).resolve().parents[2]/'pyproject.toml'
    if pyproject.exists():
        return tomllib.loads(pyproject.read_text(encoding='utf-8'))['project']['version']
    try: return version('pmo-studio')
    except PackageNotFoundError: return '0.0.0'

def _write_manifest_md(path: Path, m: DeliveryManifest) -> None:
    lines=[f'# Delivery Manifest — {m.project_slug}','',f'- Language: **{m.language}**',f'- PMO Studio version: **{m.pmo_studio_version}**',f'- Domain: **{m.domain_id} — {m.domain_label}**',f'- Model hash: `{m.model_hash}`',f'- Generated at: {m.generated_at}',f'- Zip: `{m.zip_path or "not generated"}`','', '| Kind | File | SHA256 | Size |','|---|---|---|---|']
    for a in m.artifacts: lines.append(f'| {a.kind} | {a.delivery_path} | `{a.sha256}` | {a.size_bytes} |')
    path.write_text('\n'.join(lines)+'\n', encoding='utf-8')

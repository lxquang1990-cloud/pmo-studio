from __future__ import annotations
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from docx.shared import Pt
from pmo_studio.model.builder import load_ba_model
from pmo_studio.model.schema import BAModel
from pmo_studio.model.translation import model_to_language
from pmo_studio.model.content_polish import polish_model
from pmo_studio.generators.quotation import QuotationInput, HangMuc, SubSystem, Feature as QFeature, ScreenRow, OutOfScreenItem, Assumption, generate_quotation_xlsx, MANDAY_RATE_VND

def ensure_model(project):
    path=project.root/'artifacts/model/ba-model.json'
    if not path.exists():
        from pmo_studio.model.builder import write_ba_model
        write_ba_model(project)
    return load_ba_model(project.root)

def generate_srs_from_model(project, lang='vi')->Path:
    m=polish_model(model_to_language(ensure_model(project), lang), lang); out=project.root/f'artifacts/ba/03-srs/srs-customer-ready.{lang}.md'; out.parent.mkdir(parents=True,exist_ok=True)
    title = 'SRS Customer-ready' if lang == 'en' else 'SRS Customer-ready'
    intro = 'Introduction' if lang == 'en' else 'Giới thiệu'
    scope = 'Scope' if lang == 'en' else 'Phạm vi'
    roles_label = 'User Roles' if lang == 'en' else 'Vai trò người dùng'
    func_label = 'Functional Requirements' if lang == 'en' else 'Yêu cầu chức năng'
    lines=[f'# {title} — {m.project.product or m.project.slug}','',f'## 1. {intro}',('This Software Requirements Specification is generated from the Canonical BA Model and applies to domain '+m.domain.label+'.') if lang=='en' else ('Tài liệu mô tả yêu cầu phần mềm được sinh từ Canonical BA Model, áp dụng cho domain '+m.domain.label+'.'),'', f'## 2. {scope}','### In scope', *[f'- {c.name}' for c in m.capabilities], '', '### Out of scope', *[f'- {x}' for x in m.out_of_scope], '', f'## 3. {roles_label}', *[f'- **{r.name}**: {r.description or "Vai trò nghiệp vụ trong hệ thống."}' for r in m.roles], '', f'## 4. {func_label}']
    for cap in m.capabilities:
        lines += ['', f'### {cap.id} — {cap.name}']
        for req in [r for r in m.requirements if any(f.id==r.feature_id and f.capability_id==cap.id for f in m.features)]:
            lines += ['', f'#### {req.id} — {req.name}', req.description, '', ('**User Roles:** ' if lang=='en' else '**Vai trò sử dụng:** ')+', '.join(req.actor_roles), ('**Inputs:** ' if lang=='en' else '**Đầu vào:** ')+', '.join(req.inputs), ('**Process:** ' if lang=='en' else '**Xử lý:** ')+req.process, ('**Outputs:** ' if lang=='en' else '**Đầu ra:** ')+', '.join(req.outputs), '**Business Rules:**', *[f'- {x}' for x in req.business_rules], '**Validation:**', *[f'- {x}' for x in req.validation_rules], '**Permission:**', *[f'- {x}' for x in req.permission_rules], '**Audit:**', *[f'- {x}' for x in req.audit_events], '**Acceptance Criteria:**']
            for ac in [a for a in m.acceptance_criteria if a.requirement_id==req.id]: lines.append(f'- {ac.id}: Given {ac.given}, When {ac.when}, Then {ac.then}.')
    lines += ['', '## 5. Integrations' if lang=='en' else '## 5. Tích hợp', *[f'- {x}' for x in m.integrations], '', '## 6. Non-functional Requirements', '- RBAC/permission theo vai trò.', '- Audit log cho thao tác quan trọng.', '- Import/export có kiểm soát quyền và dữ liệu.', '- Dữ liệu cần backup/restore theo chính sách vận hành.', '', '## 7. Assumptions', *[f'- {x}' for x in m.assumptions], '', '## 8. Traceability Matrix', '| REQ | Feature | Capability |', '|---|---|---|']
    for req in m.requirements:
        feat=next(f for f in m.features if f.id==req.feature_id); cap=next(c for c in m.capabilities if c.id==feat.capability_id); lines.append(f'| {req.id} | {feat.id} {feat.name} | {cap.id} {cap.name} |')
    out.write_text('\n'.join(lines)+'\n',encoding='utf-8'); _docx(out, project.root/f'exports/customer/{lang}/srs-customer-ready.docx'); return out

def generate_stories_from_model(project, lang='vi')->Path:
    m=polish_model(model_to_language(ensure_model(project), lang), lang); out=project.root/f'artifacts/ba/04-us/user-stories.{lang}.md'; out.parent.mkdir(parents=True,exist_ok=True)
    lines=[f'# User Stories & Acceptance Criteria — {m.project.product or m.project.slug}','']
    for us in m.user_stories:
        feat=next(f for f in m.features if f.id==us.feature_id)
        lines += [f'## {us.id} — {feat.name}','',f'As a **{us.role}**,',f'I want to **{us.want}**,',f'so that **{us.benefit}**.','',f'Priority: **{us.priority}**','', 'Linked Requirements:', *[f'- {r}' for r in us.linked_requirements], '', 'Acceptance Criteria:']
        for ac in [a for a in m.acceptance_criteria if a.user_story_id==us.id]: lines += [f'- **{ac.id}**', f'  - Given {ac.given}', f'  - When {ac.when}', f'  - Then {ac.then}']
        lines.append('')
    out.write_text('\n'.join(lines),encoding='utf-8'); _docx(out, project.root/f'exports/customer/{lang}/user-stories.docx'); return out

def generate_uat_from_model(project, lang='vi')->Path:
    m=polish_model(model_to_language(ensure_model(project), lang), lang); out=project.root/f'exports/customer/{lang}/uat-pack/test-cases-uat.{lang}.xlsx'; out.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook(); ws=wb.active; ws.title='Test Cases'; headers=['TC ID','Module/Capability','Linked REQ','Linked US','Linked AC','Role','Priority','Precondition','Test Data','Steps','Expected Result','Actual Result','Status','Evidence','Owner']; _header(ws,headers)
    for tc in m.test_cases:
        req=next(r for r in m.requirements if r.id==tc.requirement_id); feat=next(f for f in m.features if f.id==req.feature_id); cap=next(c for c in m.capabilities if c.id==feat.capability_id)
        ws.append([tc.id,cap.name,tc.requirement_id,tc.user_story_id,tc.acceptance_criteria_id,tc.role,tc.priority,tc.precondition,tc.test_data,'\n'.join(tc.steps),tc.expected_result,'','Not Run','', 'Customer UAT'])
    td=wb.create_sheet('Test Data'); _header(td,['Entity','Sample Data','Notes']); [td.append([x,'Dữ liệu mẫu cần chuẩn bị','']) for x in m.data_entities]
    sg=wb.create_sheet('UAT Signoff'); _header(sg,['Item','Value']); [sg.append(x) for x in [('Project',m.project.slug),('UAT Window','TBD by customer'),('Entry Criteria','No high quality findings'),('Exit Criteria','All P0/P1 passed or waived'),('Signoff','Pending')]]
    tr=wb.create_sheet('Traceability'); _header(tr,['REQ','US','AC','TC']); [tr.append([tc.requirement_id,tc.user_story_id,tc.acceptance_criteria_id,tc.id]) for tc in m.test_cases]
    df=wb.create_sheet('Defect Log Template'); _header(df,['Defect ID','TC ID','Severity','Description','Owner','Status','Evidence'])
    for sh in wb.worksheets: _autosize(sh)
    wb.save(out); return out

def generate_quote_from_model(project, lang='vi')->Path:
    m=polish_model(model_to_language(ensure_model(project), lang), lang); out=project.root/f'exports/customer/{lang}/quotation-customer-ready.{lang}.xlsx'; out.parent.mkdir(parents=True,exist_ok=True)
    subs=[]
    for cap in m.capabilities:
        qfs=[]
        for feat in [f for f in m.features if f.capability_id==cap.id]:
            items=[e for e in m.estimate_items if e.feature_id==feat.id]; md=sum(e.manday for e in items) or 3.0
            qfs.append(QFeature(feat.name, [ScreenRow(feat.name, md, note=feat.description)]))
        subs.append(SubSystem(cap.name, qfs))
    inp=QuotationInput(m.project.product or m.project.slug,[HangMuc('I. PHẦN MỀM',subs)],[OutOfScreenItem('Workshop chi tiết & baseline scope',5),OutOfScreenItem('Thiết lập dự án & DevOps',4),OutOfScreenItem('UAT, tài liệu hóa & go-live',8)],[Assumption('Phạm vi',x) for x in m.assumptions]+[Assumption('Ngoài phạm vi',x) for x in m.out_of_scope],'web','detailed',MANDAY_RATE_VND)
    return generate_quotation_xlsx(inp,out)

def export_customer_from_model(project, lang='vi'):
    return {'srs':str(generate_srs_from_model(project,lang)),'stories':str(generate_stories_from_model(project,lang)),'uat':str(generate_uat_from_model(project,lang)),'quotation':str(generate_quote_from_model(project,lang))}

def _docx(md_path:Path,out:Path):
    out.parent.mkdir(parents=True,exist_ok=True); doc=Document(); doc.styles['Normal'].font.name='Arial'; doc.styles['Normal'].font.size=Pt(10)
    for line in md_path.read_text(encoding='utf-8').splitlines():
        if line.startswith('# '): doc.add_heading(line[2:],0)
        elif line.startswith('## '): doc.add_heading(line[3:],1)
        elif line.startswith('### '): doc.add_heading(line[4:],2)
        elif line.startswith('#### '): doc.add_heading(line[5:],3)
        elif line.startswith('- '): doc.add_paragraph(line[2:],style='List Bullet')
        else: doc.add_paragraph(line)
    doc.save(out)
def _header(ws,headers):
    ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='305496'); c.alignment=Alignment(horizontal='center',wrap_text=True)
def _autosize(ws):
    thin=Side(style='thin',color='CCCCCC')
    for row in ws.iter_rows():
        for c in row: c.border=Border(left=thin,right=thin,top=thin,bottom=thin); c.alignment=Alignment(wrap_text=True,vertical='top')
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=max(12,min(50,max(len(str(c.value or '')) for c in col)+2))

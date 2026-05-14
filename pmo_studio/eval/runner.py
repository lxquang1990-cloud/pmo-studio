"""Eval and benchmark runner for PMO Studio."""
from __future__ import annotations

import json
import tempfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from statistics import mean
from typing import Any

from pmo_studio.core.project import Project
from pmo_studio.generators.stage0 import generate_stage0
from pmo_studio.generators.source_ba import generate_ba_from_sources
from pmo_studio.generators.po_pm_ic import generate_po, generate_pm, generate_ic
from pmo_studio.gates.gate_a import run_gate_a
from pmo_studio.gates.gate_bc import run_gate_b_or_c
from pmo_studio.gates.runner import run_all_gates
from pmo_studio.traceability.engine import TraceabilityEngine
from pmo_studio.traceability.validator import validate_traceability
from pmo_studio.core.manifest import create_baseline
from pmo_studio.exporters.static_html import export_static
from pmo_studio.exporters.docx_export import export_docx
from pmo_studio.metrics.recorder import MetricsRecorder
from pmo_studio.llm.factory import build_llm
from pmo_studio.llm.reviewer import build_gate_reviewer


@dataclass
class EvalCase:
    name: str
    slug: str
    customer: str
    products: list[str]
    brief: str
    source: str
    expected_id_types: set[str]
    min_trace_edges: int = 10
    min_manifest_items: int = 12
    expected_manday_min: float = 1.0
    expected_manday_max: float = 80.0


@dataclass
class EvalResult:
    name: str
    passed: bool
    details: dict[str, Any]


@dataclass
class BenchmarkResult:
    name: str
    passed: bool
    score: float
    case_results: list[EvalResult] = field(default_factory=list)
    report_path: str | None = None
    results_path: str | None = None


EVAL_CASES = [
    EvalCase(
        name="eOffice Document Workflow",
        slug="eval-eoffice",
        customer="BTECO eOffice Eval",
        products=["eOffice"],
        brief="Đánh giá PMO Studio cho dự án eOffice quản lý văn bản và phê duyệt.",
        source="""# eOffice Source
Khách hàng cần hệ thống quản lý công văn đến/đi, luồng duyệt đa cấp, phân quyền phòng ban, dashboard SLA, báo cáo quá hạn.
Tích hợp LDAP/AD, email notification, API đồng bộ văn bản với hệ thống DMS.
API_KEY=should_be_redacted_eoffice
Email liên hệ: user@example.com
""",
        expected_id_types={"SRC", "BR", "REQ", "SCR", "API", "WF", "US", "AC", "TC", "EST"},
    ),
    EvalCase(
        name="Digital Signature Platform",
        slug="eval-digital-signature",
        customer="BTECO Digital Signature Eval",
        products=["Digital Signature"],
        brief="Đánh giá PMO Studio cho nền tảng ký số và audit chứng thư.",
        source="""# Digital Signature Source
Nền tảng ký số cần quản lý certificate, signing request, ký hàng loạt PDF, audit trail, timestamp authority, callback API.
Có vai trò requester, approver, signer, auditor. Cần báo cáo số lượng hồ sơ ký theo trạng thái.
password=should_be_redacted_signature
SĐT hỗ trợ: 0912345678
""",
        expected_id_types={"SRC", "BR", "REQ", "SCR", "API", "WF", "US", "AC", "TC", "EST"},
    ),
    EvalCase(
        name="HSE Incident Management",
        slug="eval-hse",
        customer="BTECO HSE Eval",
        products=["HSE"],
        brief="Đánh giá PMO Studio cho hệ thống HSE incident/inspection/training.",
        source="""# HSE Source
Hệ thống HSE cần ghi nhận incident, near-miss, inspection checklist, corrective action, training matrix, dashboard TRIR/LTIFR.
Có workflow điều tra sự cố, phân công action owner, nhắc hạn và báo cáo compliance theo nhà máy.
token=should_be_redacted_hse
""",
        expected_id_types={"SRC", "BR", "REQ", "SCR", "API", "WF", "US", "AC", "TC", "EST"},
    ),
]


V12_SAMPLE_CASES = [
    EvalCase(
        name="Asset Management v1.2 E2E", slug="bench-asset", customer="Demo Asset", products=["Asset Management"],
        brief="Validate Asset Management pipeline.",
        source="""# Asset Management
| STT | CHỨC NĂNG | MÔ TẢ |
| 1 | Quản lý Danh mục Tài sản / TTB | Mã tài sản, serial, phòng ban, người sử dụng |
| 2 | Cấp phát bàn giao thu hồi điều chuyển | Workflow tài sản, audit, evidence |
| 3 | Kiểm kê bảo trì thanh lý | Campaign kiểm kê, bảo trì, báo cáo tài sản |
""",
        expected_id_types={"SRC", "BR", "REQ", "SCR", "API", "WF", "US", "AC", "TC", "EST"}, min_trace_edges=50, expected_manday_max=180.0
    ),
    EvalCase(
        name="LegalIQ v1.2 E2E", slug="bench-legaliq", customer="PVCFC Legal", products=["LegalIQ"],
        brief="Validate LegalIQ pipeline.",
        source="""# LegalIQ
| STT | CHỨC NĂNG | MÔ TẢ |
| 1 | Hỏi đáp pháp lý | FAQ, AI trả lời có trích dẫn nguồn |
| 2 | Quản lý ủy quyền | Trình ký eOffice, ban hành, đồng bộ |
| 3 | Hợp đồng thẩm định | PMS integration, báo cáo pháp lý, RBAC audit |
""",
        expected_id_types={"SRC", "BR", "REQ", "SCR", "API", "WF", "US", "AC", "TC", "EST"}, min_trace_edges=50, expected_manday_max=220.0
    ),
    EvalCase(
        name="CRM v1.2 E2E", slug="bench-crm", customer="Demo CRM", products=["CRM"],
        brief="Validate CRM source-driven pipeline.",
        source="""# CRM
| STT | CHỨC NĂNG | MÔ TẢ |
| 1 | Quản lý Lead | Capture, qualify, assign lead |
| 2 | Opportunity Pipeline | Stage, forecast, approval workflow |
| 3 | Sales Dashboard | Báo cáo doanh số, conversion, export XLSX |
""",
        expected_id_types={"SRC", "BR", "REQ", "SCR", "API", "WF", "US", "AC", "TC", "EST"}, min_trace_edges=50, expected_manday_max=180.0
    ),
]

def run_eval(root: Path | None = None) -> EvalResult:
    """Backward-compatible single eval used by earlier phases."""
    bench = run_benchmark(root=root, cases=[EVAL_CASES[0]], export_docx_enabled=False)
    result = bench.case_results[0]
    return EvalResult("eval-bteco-doc", result.passed, result.details)


def run_benchmark(root: Path | None = None, cases: list[EvalCase] | None = None, export_docx_enabled: bool = True, llm_provider: str = "noop", llm_model: str | None = None, gate_reviewer_provider: str | None = None, gate_reviewer_model: str | None = None, gate_cache_root: str | Path | None = None, gate_timeout: int | None = None, gate_fallback: bool = False) -> BenchmarkResult:
    temp_ctx = tempfile.TemporaryDirectory() if root is None else None
    root_path = root or Path(temp_ctx.name)
    root_path.mkdir(parents=True, exist_ok=True)
    selected = cases or EVAL_CASES
    results = [_run_case(root_path, case, export_docx_enabled=export_docx_enabled, llm_provider=llm_provider, llm_model=llm_model, gate_reviewer_provider=gate_reviewer_provider, gate_reviewer_model=gate_reviewer_model, gate_cache_root=gate_cache_root, gate_timeout=gate_timeout, gate_fallback=gate_fallback) for case in selected]
    scores = [float(r.details.get("score", 0.0)) for r in results]
    bench = BenchmarkResult(
        name="pmo-studio-benchmark",
        passed=all(r.passed for r in results),
        score=round(mean(scores), 4) if scores else 0.0,
        case_results=results,
    )
    out_dir = root_path / "benchmark"
    out_dir.mkdir(parents=True, exist_ok=True)
    results_path = out_dir / "results.json"
    report_path = out_dir / "report.md"
    payload = asdict(bench)
    payload["results_path"] = str(results_path)
    payload["report_path"] = str(report_path)
    results_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report_path.write_text(_benchmark_report(bench), encoding="utf-8")
    bench.results_path = str(results_path)
    bench.report_path = str(report_path)
    return bench


def _run_case(root_path: Path, case: EvalCase, export_docx_enabled: bool = True, llm_provider: str = "noop", llm_model: str | None = None, gate_reviewer_provider: str | None = None, gate_reviewer_model: str | None = None, gate_cache_root: str | Path | None = None, gate_timeout: int | None = None, gate_fallback: bool = False) -> EvalResult:
    source = root_path / f"{case.slug}-source.md"
    source.write_text(case.source, encoding="utf-8")
    p = Project.create(case.slug, customer=case.customer, root_base=root_path)
    recorder = MetricsRecorder(p.root, run_id="eval")
    llm = build_llm(llm_provider, llm_model) if llm_provider and llm_provider != "noop" else None
    with recorder.span("eval.case", case=case.name, llm_provider=llm_provider):
        generate_stage0(p, brief=case.brief, products=case.products, sources=[source])
        generate_po(p); generate_pm(p); generate_ba_from_sources(p, llm=llm); generate_ic(p)
        TraceabilityEngine(p.root).write_outputs()
        trace = validate_traceability(p.root)
        reviewer_provider = gate_reviewer_provider if gate_reviewer_provider is not None else llm_provider
        reviewer_model = gate_reviewer_model if gate_reviewer_model is not None else llm_model
        gate_reviewer = build_gate_reviewer(reviewer_provider, reviewer_model, cache_root=gate_cache_root, fallback_on_error=gate_fallback, timeout=gate_timeout) if reviewer_provider and reviewer_provider != "noop" else None
        quality = run_all_gates(p, include_c=False, reviewer=gate_reviewer)
        manifest = create_baseline(p.root, p.config.project_slug, "v1.0")
        html = export_static(p.root)
        docx = export_docx(p.root) if export_docx_enabled else None
    recorder.save()
    graph = json.loads((p.root / "traceability" / "graph.json").read_text(encoding="utf-8"))
    found_types = {n.get("type") for n in graph.get("nodes", [])}
    missing_types = sorted(case.expected_id_types - found_types)
    redacted_ok = _redacted_ok(p.root)
    quotation = _quotation_sanity(p.root, case)
    checks = {
        "gate_pass_rate": quality["passed"] / max(quality["total"], 1),
        "trace_passed": trace.passed,
        "trace_edges": trace.edge_count,
        "trace_edge_min_passed": trace.edge_count >= case.min_trace_edges,
        "missing_id_types": missing_types,
        "id_coverage_passed": not missing_types,
        "redacted_ok": redacted_ok,
        "manifest_items": manifest.item_count,
        "manifest_min_passed": manifest.item_count >= case.min_manifest_items,
        "dashboard_exists": html.exists(),
        "docx_exists": bool(docx.exists()) if docx else True,
        "quotation_total_manday": quotation["total_manday"],
        "quotation_sane": quotation["sane"],
        "project_root": str(p.root),
    }
    score = _score(checks)
    checks["score"] = score
    passed = score >= 0.9 and checks["trace_passed"] and checks["redacted_ok"] and checks["quotation_sane"]
    return EvalResult(case.name, passed, checks)


def _redacted_ok(project_root: Path) -> bool:
    bad = ["should_be_redacted", "user@example.com", "0912345678"]
    for path in (project_root / "source" / "redacted").glob("*"):
        text = path.read_text(encoding="utf-8", errors="ignore")
        if any(token in text for token in bad):
            return False
    return True


def _quotation_sanity(project_root: Path, case: EvalCase) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(project_root / "artifacts" / "ba" / "06-quotation.xlsx", data_only=True)
        total = 0.0
        if "Summary" in wb.sheetnames:
            ws = wb["Summary"]
            metrics = {str(row[0].value): row[1].value for row in ws.iter_rows(min_row=2) if row[0].value is not None}
            total = float(metrics.get("Total manday", 0) or 0)
        elif "Tổng hợp" in wb.sheetnames:
            ws = wb["Tổng hợp"]
            for row in ws.iter_rows(values_only=True):
                label = str(row[1] or "") if len(row) > 1 else ""
                if "Tổng manday cuối" in label:
                    total = float(row[2] or 0)
                    break
        return {"total_manday": total, "sane": case.expected_manday_min <= total <= case.expected_manday_max}
    except Exception as exc:
        return {"total_manday": 0.0, "sane": False, "error": str(exc)}


def _score(checks: dict[str, Any]) -> float:
    weighted = [
        (checks["gate_pass_rate"], 0.25),
        (1.0 if checks["trace_passed"] else 0.0, 0.2),
        (1.0 if checks["trace_edge_min_passed"] else 0.1, 0.1),
        (1.0 if checks["id_coverage_passed"] else 0.0, 0.15),
        (1.0 if checks["redacted_ok"] else 0.0, 0.15),
        (1.0 if checks["manifest_min_passed"] else 0.0, 0.05),
        (1.0 if checks["dashboard_exists"] and checks["docx_exists"] else 0.0, 0.05),
        (1.0 if checks["quotation_sane"] else 0.0, 0.05),
    ]
    return round(sum(value * weight for value, weight in weighted), 4)


def _benchmark_report(bench: BenchmarkResult) -> str:
    lines = [
        "# PMO Studio Benchmark Report",
        "",
        f"**Benchmark:** {bench.name}",
        f"**Status:** {'PASS' if bench.passed else 'FAIL'}",
        f"**Average score:** {bench.score:.2%}",
        "",
        "| Case | Status | Score | Gate Pass Rate | Trace Edges | Redaction | Quotation md |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for result in bench.case_results:
        d = result.details
        lines.append(
            f"| {result.name} | {'PASS' if result.passed else 'FAIL'} | {float(d['score']):.2%} | "
            f"{float(d['gate_pass_rate']):.2%} | {d['trace_edges']} | {'OK' if d['redacted_ok'] else 'FAIL'} | {d['quotation_total_manday']} |"
        )
    lines.extend(["", "## Case Details", ""])
    for result in bench.case_results:
        lines.append(f"### {result.name}")
        lines.append(f"- Status: {'PASS' if result.passed else 'FAIL'}")
        lines.append(f"- Score: {float(result.details['score']):.2%}")
        lines.append(f"- Project root: `{result.details['project_root']}`")
        if result.details.get("missing_id_types"):
            lines.append(f"- Missing ID types: {', '.join(result.details['missing_id_types'])}")
        lines.append("")
    return "\n".join(lines) + "\n"


def main() -> None:
    result = run_benchmark()
    print(json.dumps(asdict(result), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

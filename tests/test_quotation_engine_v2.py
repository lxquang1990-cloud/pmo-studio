from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from pmo_studio.generators.quotation import generic_quotation_input, generate_quotation_xlsx


def _joined_workbook(path: Path) -> str:
    wb = load_workbook(path, data_only=False)
    return "\n".join(
        str(c.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    ).lower()


def test_quotation_v2_has_role_breakdown_and_rationale_sheets(tmp_path):
    source = """
| STT | CHỨC NĂNG | MÔ TẢ |
| 1 | Quản lý Lead | Capture leads and assign to sales |
| 2 | AI Sales Forecast | AI forecast with CRM API integration and dashboard export |
| 3 | Discount Approval Workflow | Phê duyệt discount theo SLA, RBAC and audit |
"""
    out = tmp_path / "quotation-v2.xlsx"
    generate_quotation_xlsx(generic_quotation_input("CRM", source_text=source), out)
    wb = load_workbook(out, data_only=False)
    assert "Role Breakdown" in wb.sheetnames
    assert "Estimate Rationale" in wb.sheetnames
    assert "Estimate Detail" not in wb.sheetnames  # added by project wrapper only
    text = _joined_workbook(out)
    assert "ai sales forecast" in text
    assert "role breakdown" not in text  # sheet name is structural, not cell data
    assert "complexity=very_complex" in text or "complexity=complex" in text
    assert "roles={" in text
    assert "risk buffer" in text
    assert "quản lý danh mục tài sản" not in text
    assert "legaliq" not in text


def test_project_wrapper_appends_v2_estimate_detail(tmp_path):
    from pmo_studio.core.project import Project, ProjectConfig
    from pmo_studio.generators.quotation import generate_quotation_for_project

    project = Project(tmp_path / "crm", ProjectConfig(project_slug="crm", customer="Demo"))
    project.ensure_layout(); project.save()
    (project.root / "source/redacted/source.md").write_text("CRM Lead Opportunity API workflow dashboard", encoding="utf-8")
    out = generate_quotation_for_project(project)
    wb = load_workbook(out, data_only=False)
    assert "Estimate Detail" in wb.sheetnames
    ws = wb["Estimate Detail"]
    headers = [c.value for c in ws[1]]
    assert headers[:7] == ["EST ID", "Module", "Function", "Work Item Type", "Work Item ID", "Complexity", "Rationale"]
    assert "BA" in headers and "Dev" in headers and "QA" in headers
    assert ws.max_row >= 2

from pathlib import Path
from openpyxl import load_workbook

from pmo_studio.decomposition.engine import decompose_project, write_decomposition, load_decomposition


def test_v30_functional_decomposition_outputs_json_md_xlsx(tmp_path):
    root = tmp_path / "asset-demo"
    (root / "source" / "redacted").mkdir(parents=True)
    (root / "source" / "redacted" / "brief.md").write_text("Quản lý tài sản gồm cấp phát bàn giao kiểm kê bảo trì thanh lý báo cáo", encoding="utf-8")
    out = write_decomposition(root)
    assert out.exists()
    assert (root / "artifacts" / "ba" / "00-functional-decomposition.md").exists()
    xlsx = root / "artifacts" / "ba" / "00-functional-decomposition.xlsx"
    assert xlsx.exists()
    wb = load_workbook(xlsx)
    assert "Capability Map" in wb.sheetnames
    assert "Work Items" in wb.sheetnames
    report = load_decomposition(root)
    assert report is not None
    assert report.capabilities
    assert any(f.work_items for c in report.capabilities for f in c.features)


def test_v30_decomposition_has_quality_score(tmp_path):
    root = tmp_path / "generic-demo"
    (root / "source" / "redacted").mkdir(parents=True)
    (root / "source" / "redacted" / "brief.md").write_text("Core workflow report permission audit", encoding="utf-8")
    report = decompose_project(root)
    assert report.schema == "pmo.functional_decomposition.v1"
    assert 0 <= report.score <= 100
    assert all(f.completeness_score > 0 for c in report.capabilities for f in c.features)

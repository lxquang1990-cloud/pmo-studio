from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from pmo_studio.core.project import Project, ProjectConfig
from pmo_studio.generators.po_pm_ic import generate_po, generate_pm
from pmo_studio.generators.source_ba import generate_ba_from_sources


def test_legaliq_does_not_fall_back_to_asset_management():
    with TemporaryDirectory() as td:
        root_base = Path(td)
        project_root = root_base / "legal-regression"
        project_root.mkdir(parents=True)
        source_dir = project_root / "source" / "redacted"
        source_dir.mkdir(parents=True)
        (source_dir / "source.md").write_text(
            "LIQ LegalIQ Phần mềm Hỏi đáp Pháp lý. "
            "Bao gồm FAQ, AI trả lời có trích dẫn, B.PCTT escalation, "
            "quản lý ủy quyền, hợp đồng, thẩm định, eOffice/PMS.",
            encoding="utf-8",
        )
        project = Project(project_root, ProjectConfig(project_slug="legal-regression", customer="PVCFC / Legal Department"))

        generate_po(project)
        generate_pm(project)
        generate_ba_from_sources(project)

        text = "\n".join(
            (project_root / rel).read_text(encoding="utf-8")
            for rel in [
                "artifacts/po/01-vision.md",
                "artifacts/pm/01-charter.md",
                "artifacts/ba/01-prd.md",
                "artifacts/ba/02-brd.md",
                "artifacts/ba/05-test-cases.md",
            ]
        ).lower()
        assert "legaliq" in text
        assert "pháp lý" in text
        assert "asset management" not in text
        assert "quản lý trang thiết bị" not in text
        assert "ttb/tài sản" not in text

        wb = load_workbook(project_root / "artifacts/ba/06-quotation.xlsx", data_only=True)
        quote = "\n".join(
            str(cell)
            for ws in wb.worksheets
            for row in ws.iter_rows(values_only=True)
            for cell in row
            if cell is not None
        ).lower()
        assert "hỏi & đáp pháp lý" in quote
        assert "ủy quyền" in quote
        assert "quản lý danh mục tài sản" not in quote
        assert "asset management" not in quote

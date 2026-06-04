from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from pmo_studio.webdoc.ingest import ingest_webdoc_discovery, load_browser_capture, build_test_cases


def test_webdoc_ingest_generates_pmo_artifacts(tmp_path: Path):
    discovery = tmp_path / "discovery"
    discovery.mkdir()
    (discovery / "registration.json").write_text(json.dumps({
        "route": "registration",
        "url": "https://example.test/app#/registration",
        "title": "PMS",
        "headers": ["Mã", "Tiêu đề", "Năm KH", "Tình trạng"],
        "rows": ["26-KTNB-01 2026 Đang xử lý"],
        "bodyText": "Kế hoạch năm Đang xử lý Đã duyệt",
    }, ensure_ascii=False), encoding="utf-8")
    (discovery / "request-purchase.json").write_text(json.dumps({
        "route": "request-purchase",
        "url": "https://example.test/app#/request-purchase",
        "headers": ["Số YCMS", "Đơn vị yêu cầu", "Tình trạng", "Ngày yêu cầu"],
        "rows": ["0024/26/YCMS-TM Ban Thương mại Đã duyệt"],
        "bodyText": "Yêu cầu mua sắm",
    }, ensure_ascii=False), encoding="utf-8")
    project_root = tmp_path / "pmo-project"

    result = ingest_webdoc_discovery(project_root, discovery)

    assert result.screens == 2
    assert result.test_cases >= 8  # 2 screen-derived + guardrail cases
    assert result.markdown.exists()
    assert result.workbook.exists()
    assert result.discovery_summary.exists()
    text = result.markdown.read_text(encoding="utf-8")
    assert "WEB-AI-STAT-MAP-001" in text
    assert "Registration" in text
    wb = load_workbook(result.workbook, read_only=True)
    assert wb.sheetnames == ["00_Bia", "00_TomTat_KetQua_Test", "01_ChuThich", "02_TatCa_TestCases"]
    assert wb["02_TatCa_TestCases"].max_row == result.test_cases + 3
    headers = [wb["02_TatCa_TestCases"].cell(3, c).value for c in range(1, 16)]
    assert headers == ["Mã TC", "Module", "Chức năng", "Loại test", "Priority", "Tiêu đề testcase", "Tiền điều kiện", "Các bước thực hiện", "Dữ liệu test", "Kết quả mong đợi", "Kết quả thực tế", "Trạng thái", "Tester", "Ngày test", "Ghi chú"]


def test_webdoc_capture_infers_modules(tmp_path: Path):
    (tmp_path / "order-plan.json").write_text(json.dumps({
        "route": "request-order-plan",
        "headers": ["Số ĐH kế hoạch", "Giá trị khái toán", "Tiền tệ"],
        "bodyText": "Bảng theo dõi đơn hàng kế hoạch",
    }, ensure_ascii=False), encoding="utf-8")
    screens = load_browser_capture(tmp_path)
    cases = build_test_cases(screens)
    assert screens[0].module == "RequestOrder"
    assert any("Giá trị khái toán" in row[11] for row in cases)

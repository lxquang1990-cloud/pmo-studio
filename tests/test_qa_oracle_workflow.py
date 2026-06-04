from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from pmo_studio.qa_oracle.config import load_adapter
from pmo_studio.qa_oracle.snapshot import build_snapshot_from_json, write_snapshot
from pmo_studio.qa_oracle.testcases import build_oracle_testcases, export_oracle_workbook
from pmo_studio.qa_oracle.quality import check_oracle_workbook


def test_oracle_snapshot_and_workbook_quality(tmp_path: Path):
    adapter_path = Path('pmo_studio/domain/packs/qa_oracle/pms_procurement.yaml')
    adapter = load_adapter(adapter_path)
    raw = {
        'results': {'QueueAI': {'fetched': 1646}},
        'objectType': {'Registration': 2, 'RequestPurchase': 931, 'RequestOrder': 713},
    }
    snapshot = build_snapshot_from_json(adapter, raw, snapshot_at='2026-06-04T16:09:29+07:00')
    assert len(snapshot.metrics) == 4
    assert {m.metric_id: m.value for m in snapshot.metrics}['request_purchase_ai_synced_count'] == 931

    snap_path = write_snapshot(snapshot, tmp_path / 'oracle_snapshot.json')
    assert json.loads(snap_path.read_text())['metrics'][1]['value'] == 931

    cases = build_oracle_testcases(adapter, snapshot)
    assert len(cases) == 10  # 7 question templates + 3 guardrails
    assert 'Số liệu oracle kỳ vọng: 931' in cases[2].ket_qua_mong_doi

    workbook = export_oracle_workbook(cases, snapshot, tmp_path / 'oracle_testcases.xlsx', project='PMS', module='Procurement')
    wb = load_workbook(workbook, data_only=True)
    assert '02_TatCa_TestCases' in wb.sheetnames
    assert '03_Oracle_Snapshot' in wb.sheetnames
    quality = check_oracle_workbook(workbook, min_cases=10)
    assert quality.passed, quality.to_dict()
    assert quality.counts['test_cases'] == 10


def test_quality_gate_blocks_stale_no_snapshot_text(tmp_path: Path):
    adapter = load_adapter('pmo_studio/domain/packs/qa_oracle/pms_procurement.yaml')
    raw = {'results': {'QueueAI': {'fetched': 1}}, 'objectType': {'Registration': 1, 'RequestPurchase': 1, 'RequestOrder': 1}}
    snapshot = build_snapshot_from_json(adapter, raw)
    cases = build_oracle_testcases(adapter, snapshot)
    cases[0].ket_qua_mong_doi = 'Chưa có snapshot dữ liệu realtime độc lập.'
    workbook = export_oracle_workbook(cases, snapshot, tmp_path / 'bad.xlsx', project='PMS', module='Procurement')
    quality = check_oracle_workbook(workbook, min_cases=1)
    assert not quality.passed
    assert any(i.code == 'STALE_NO_SNAPSHOT_TEXT' for i in quality.issues)

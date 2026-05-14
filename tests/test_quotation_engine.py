from openpyxl import load_workbook

from pmo_studio.generators.quotation import generic_quotation_input, generate_quotation_xlsx


def test_generic_quotation_reads_source_features_and_complexity(tmp_path):
    source = """
| STT | CHỨC NĂNG | MÔ TẢ |
| 1 | Quản lý Lead | Capture leads |
| 2 | AI Sales Forecast | AI forecast with CRM API integration |
| 3 | Approval Workflow | Phê duyệt discount theo SLA |
"""
    out = tmp_path / "quotation.xlsx"
    generate_quotation_xlsx(generic_quotation_input("CRM", source_text=source), out)
    wb = load_workbook(out, data_only=False)
    joined = "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None).lower()
    assert "quản lý lead" in joined
    assert "ai sales forecast" in joined
    assert "complexity" in joined
    assert "quản lý danh mục tài sản" not in joined
    assert "legaliq" not in joined

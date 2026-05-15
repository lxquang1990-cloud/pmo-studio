from pathlib import Path
from openpyxl import Workbook, load_workbook
from pmo_studio.core.project import Project
from pmo_studio.generators.stage0 import generate_stage0
from pmo_studio.model.builder import write_ba_model
from pmo_studio.model.delivery_pack import export_delivery_pack


def _wb(path: Path, rows: int):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb=Workbook(); ws=wb.active; ws.title='Feature List'; ws.append(['STT','Chức năng','Đơn giá','Manday','Thành tiền','Ghi chú'])
    for i in range(rows): ws.append([i+1, f'Detailed feature {i}', 3900000, 1, 3900000, 'note'])
    wb.create_sheet('Tổng hợp'); wb.create_sheet('Giả định'); wb.save(path)


def test_delivery_pack_prefers_detailed_quotation(tmp_path):
    src=tmp_path/'asset.md'; src.write_text('Quản lý tài sản cấp phát kiểm kê bảo trì thanh lý', encoding='utf-8')
    p=Project.create('asset-quote-detail', customer='Demo', root_base=tmp_path)
    generate_stage0(p, brief='asset', products=['Asset'], sources=[src]); write_ba_model(p)
    detailed=p.root/'exports/customer/vi/quotation-customer-ready-detailed.vi.xlsx'
    _wb(detailed, 30)
    export_delivery_pack(p, 'vi', make_zip=False)
    delivered=p.root/'exports/delivery/vi/04-quotation-customer-ready.vi.xlsx'
    wb=load_workbook(delivered, read_only=True, data_only=True)
    assert wb['Feature List'].max_row == 31

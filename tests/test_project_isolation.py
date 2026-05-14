from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from pmo_studio.core.project import Project, ProjectConfig
from pmo_studio.generators.po_pm_ic import generate_po, generate_pm
from pmo_studio.generators.source_ba import generate_ba_from_sources


def _make_project(root_base: Path, slug: str, source: str) -> Project:
    project_root = root_base / slug
    project_root.mkdir(parents=True)
    source_dir = project_root / "source" / "redacted"
    source_dir.mkdir(parents=True)
    (source_dir / "source.md").write_text(source, encoding="utf-8")
    return Project(project_root, ProjectConfig(project_slug=slug, customer="Demo Customer"))


def _text(project: Project) -> str:
    paths = [
        "artifacts/po/01-vision.md",
        "artifacts/pm/01-charter.md",
        "artifacts/ba/01-prd.md",
        "artifacts/ba/02-brd.md",
        "artifacts/ba/05-test-cases.md",
    ]
    out = "\n".join((project.root / p).read_text(encoding="utf-8") for p in paths)
    wb = load_workbook(project.root / "artifacts/ba/06-quotation.xlsx", data_only=True)
    out += "\n" + "\n".join(
        str(cell)
        for ws in wb.worksheets
        for row in ws.iter_rows(values_only=True)
        for cell in row
        if cell is not None
    )
    return out.lower()


def test_projects_are_folder_isolated_and_do_not_cross_read_sources():
    with TemporaryDirectory() as td:
        root_base = Path(td)
        crm = _make_project(
            root_base,
            "crm-isolation",
            "| STT | CHỨC NĂNG | MÔ TẢ |\n"
            "| 1 | Quản lý Lead | Capture and qualify leads |\n"
            "| 2 | Quản lý Opportunity | Sales pipeline and forecast |\n",
        )
        legal = _make_project(
            root_base,
            "legal-isolation",
            "LIQ LegalIQ hỏi đáp pháp lý, quản lý ủy quyền, hợp đồng, B.PCTT escalation.",
        )

        for project in [crm, legal]:
            generate_po(project)
            generate_pm(project)
            generate_ba_from_sources(project)

        crm_text = _text(crm)
        legal_text = _text(legal)

        assert "quản lý lead" in crm_text
        assert "opportunity" in crm_text
        assert "legaliq" not in crm_text
        assert "ủy quyền" not in crm_text

        assert "legaliq" in legal_text
        assert "ủy quyền" in legal_text
        assert "quản lý lead" not in legal_text
        assert "opportunity" not in legal_text

        assert (crm.root / "artifacts/ba/01-prd.md").exists()
        assert (legal.root / "artifacts/ba/01-prd.md").exists()
        assert crm.root != legal.root

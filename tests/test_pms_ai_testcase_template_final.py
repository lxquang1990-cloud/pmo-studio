from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from pmo_studio.exporters.pms_ai_testcase_template import (
    CANONICAL_COLUMNS,
    REMOVED_TECHNICAL_COLUMNS,
    PmsAiTestCase,
    create_blank_pms_ai_template,
    export_pms_ai_testcases,
    extract_visible_ai_answer,
)


def test_extract_visible_ai_answer_strips_api_wrapper():
    raw = 'Executed at: 2026-06-03T22:12:50\nHTTP status: 200\nElapsed: 18.02s\nAI answer:\n{"data":{"answer":"Nội dung trả lời AI"},"success":true}'
    assert extract_visible_ai_answer(raw) == "Nội dung trả lời AI"


def test_final_template_exports_one_ba_qa_sheet(tmp_path: Path):
    out = tmp_path / "pms-ai-final.xlsx"
    export_pms_ai_testcases([
        PmsAiTestCase(
            ma_tc="PMS-AI-001",
            module="PMS AI Q&A",
            chuc_nang="Thống kê PMS",
            loai_test="FACT",
            priority="P1",
            tieu_de="Có bao nhiêu kế hoạch năm?",
            tien_dieu_kien="User có quyền xem PMS.",
            cac_buoc="Nhập câu hỏi vào AI.",
            du_lieu_test="Câu hỏi: Có bao nhiêu kế hoạch năm?",
            ket_qua_mong_doi="AI trả đúng số lượng kế hoạch năm.",
            ket_qua_thuc_te='AI answer:\n{"data":{"answer":"Có 2 kế hoạch năm."}}',
            trang_thai="PASS_AUTO",
            ghi_chu="Generated từ danh sách câu hỏi end-user thống kê PMS AI. Không dùng wording backend/máy móc; ưu tiên câu hỏi tự nhiên theo vai trò.",
            group_sheet="02_Nhom_A_Factual",
        )
    ], out)

    wb = load_workbook(out, data_only=True)
    assert wb.sheetnames == ["00_Bia", "00_TomTat_KetQua_Test", "01_ChuThich", "02_TatCa_TestCases"]
    ws = wb["02_TatCa_TestCases"]
    headers = [ws.cell(3, c).value for c in range(1, len(CANONICAL_COLUMNS) + 1)]
    assert headers == CANONICAL_COLUMNS
    for removed in REMOVED_TECHNICAL_COLUMNS:
        assert removed not in headers
    assert ws.cell(4, 4).value == "A - Dữ kiện trực tiếp"
    assert ws.cell(4, 11).value == "Có 2 kế hoạch năm."
    assert ws.cell(4, 12).value == "Pass"
    assert ws.cell(4, 15).value in (None, "")


def test_blank_template_has_required_sheets_and_headers(tmp_path: Path):
    out = tmp_path / "blank.xlsx"
    create_blank_pms_ai_template(out)
    wb = load_workbook(out, data_only=True)
    assert wb.sheetnames == ["00_Bia", "00_TomTat_KetQua_Test", "01_ChuThich", "02_TatCa_TestCases"]
    ws = wb["02_TatCa_TestCases"]
    assert [ws.cell(3, c).value for c in range(1, len(CANONICAL_COLUMNS) + 1)] == CANONICAL_COLUMNS
    assert ws.max_row == 3

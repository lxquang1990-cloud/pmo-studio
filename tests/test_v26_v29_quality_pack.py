from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook

from pmo_studio.quality.output_quality import write_output_quality
from pmo_studio.quality.source_refiner import write_source_refinement
from pmo_studio.exporters.uat_pack import export_uat_pack
from pmo_studio.quality.consistency_engine import write_consistency_report, run_consistency_engine


def _project(tmp_path: Path) -> Path:
    root = tmp_path / "demo"
    (root / "source" / "redacted").mkdir(parents=True)
    (root / "source" / "redacted" / "brief.md").write_text("RFID kiểm kê tài sản bảo trì dashboard", encoding="utf-8")
    (root / "artifacts" / "ba" / "03-srs").mkdir(parents=True)
    (root / "artifacts" / "ic").mkdir(parents=True)
    (root / "artifacts" / "ba" / "01-prd.md").write_text("# PRD\n\nTBD\n\nREQ-CORE-001\nAC-001-01\nTC-001\n", encoding="utf-8")
    (root / "artifacts" / "ba" / "02-brd.md").write_text("# BRD\n\nREQ-CORE-001\nAC-001-01\nTC-001\n", encoding="utf-8")
    (root / "artifacts" / "ba" / "03-srs" / "srs.md").write_text("# SRS\n\nREQ-CORE-001\nAC-001-01\nTC-001\n", encoding="utf-8")
    (root / "artifacts" / "ba" / "05-test-cases.md").write_text("# Test Cases\n\nTC-001 linked AC-001-01 expected result\n", encoding="utf-8")
    (root / "artifacts" / "ic" / "04-uat-plan.md").write_text("# UAT Plan\n\nPass/Fail tracking\n", encoding="utf-8")
    return root


def test_v26_output_quality_replaces_placeholders(tmp_path):
    root = _project(tmp_path)
    out = write_output_quality(root)
    assert out.exists()
    assert "TBD" not in (root / "artifacts" / "ba" / "01-prd.md").read_text(encoding="utf-8")
    assert (root / "quality" / "output-quality.md").exists()


def test_v27_source_refiner_adds_missing_source_notes(tmp_path):
    root = _project(tmp_path)
    out = write_source_refinement(root)
    assert out.exists()
    brd = (root / "artifacts" / "ba" / "02-brd.md").read_text(encoding="utf-8")
    assert "Source-grounded Coverage Notes" in brd
    assert "RFID" in brd


def test_v28_uat_pack_exports_excel_and_zip(tmp_path):
    root = _project(tmp_path)
    out = export_uat_pack(root)
    assert out.exists()
    with ZipFile(out) as z:
        assert "test-cases-uat.xlsx" in z.namelist()
    wb = load_workbook(out.parent / "test-cases-uat.xlsx")
    assert "Test Cases" in wb.sheetnames
    assert "UAT Signoff" in wb.sheetnames
    assert wb["Test Cases"]["A1"].value == "TC ID"


def test_v29_consistency_engine_reports_clean_after_quality(tmp_path):
    root = _project(tmp_path)
    write_output_quality(root)
    out = write_consistency_report(root)
    report = run_consistency_engine(root)
    assert out.exists()
    assert report.schema == "pmo.consistency_engine.v2"
    assert all(f.id != "CE-TEMPLATE" for f in report.findings)

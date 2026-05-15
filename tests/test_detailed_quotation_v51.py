from openpyxl import load_workbook
from pmo_studio.core.project import Project
from pmo_studio.generators.stage0 import generate_stage0
from pmo_studio.model.builder import write_ba_model
from pmo_studio.model.detailed_quotation import generate_detailed_quotation
from pmo_studio.model.delivery_pack import export_delivery_pack

def test_legaliq_detailed_quote_customer_ready(tmp_path):
    text='LegalIQ pháp lý hỏi đáp AI ủy quyền hợp đồng thẩm định eOffice PMS SmartCA AD Azure MFA'
    src=tmp_path/'sow.md'; src.write_text(text, encoding='utf-8')
    p=Project.create('legaliq-detail', customer='PVCFC', root_base=tmp_path)
    generate_stage0(p, brief=text, products=['LegalIQ'], sources=[src]); write_ba_model(p)
    out=generate_detailed_quotation(p,'vi')
    wb=load_workbook(out,data_only=True)
    assert wb.sheetnames == ['Feature List','Tổng hợp','Giả định']
    assert wb['Feature List'].max_row >= 35
    text='\n'.join(str(c.value or '') for row in wb['Feature List'].iter_rows() for c in row)
    assert 'AI gợi ý câu trả lời pháp lý dạng draft' in text
    assert 'Trích dẫn nguồn/citation' in text
    assert 'Cấu hình tích hợp AD/Azure' in text
    export_delivery_pack(p,'vi',make_zip=False)
    delivered=load_workbook(p.root/'exports/delivery/vi/04-quotation-customer-ready.vi.xlsx',data_only=True)
    assert delivered.sheetnames == ['Feature List','Tổng hợp','Giả định']
    assert delivered['Feature List'].max_row >= 35

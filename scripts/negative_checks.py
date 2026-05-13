#!/usr/bin/env python3
"""Negative checks for PMO Studio guardrails.

No pytest required. Exits non-zero if a guardrail does not fail as expected.
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook

from pmo_studio.core.project import Project
from pmo_studio.generators.stage0 import generate_stage0
from pmo_studio.generators.source_ba import generate_ba_from_sources
from pmo_studio.generators.po_pm_ic import generate_ic
from pmo_studio.gates.gate_a import run_gate_a
from pmo_studio.interfaces.telegram_adapter import handle_command_structured
from pmo_studio.traceability.engine import TraceabilityEngine
from pmo_studio.exporters.bundle import export_bundle


def main() -> None:
    root = Path(tempfile.mkdtemp(prefix="pmo-negative-"))
    try:
        source = root / "source.md"
        source.write_text("Khách hàng cần eOffice. token=should_be_redacted_negative", encoding="utf-8")
        p = Project.create("negative-demo", customer="Negative Customer", root_base=root)
        generate_stage0(p, brief="Negative checks", sources=[source])
        generate_ba_from_sources(p)
        generate_ic(p)
        TraceabilityEngine(p.root).write_outputs()

        invalid = p.root / "artifacts" / "ba" / "invalid-id.md"
        invalid.write_text("# SRS\n\n## 1. Introduction\nBAD-001 should fail\n\n## 2. Overall Description\nX\n\n## 3. Specific Requirements\nX\n", encoding="utf-8")
        r = run_gate_a(invalid, "ba.srs")
        assert not r.passed, "Invalid ID artifact should fail Gate A"

        quote = p.root / "artifacts" / "ba" / "06-quotation.xlsx"
        wb = load_workbook(quote)
        ws = wb["Summary"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "Total cost VND":
                row[1].value = 1
        wb.save(quote)
        r = run_gate_a(quote, "ba.quotation")
        assert not r.passed, "Quotation total mismatch should fail Gate A"

        cfg = p.root / "artifacts" / "ic" / "02-config-workbook.xlsx"
        wb = load_workbook(cfg)
        ws = wb["Org Settings"]
        ws.append(["api_token", "plaintext-secret", "yes", "IT", "Draft"])
        wb.save(cfg)
        r = run_gate_a(cfg, "ic.config_workbook")
        assert not r.passed, "Plaintext secret in config workbook should fail Gate A"

        resp = handle_command_structured("/pmo baseline negative-demo v1.0", root=root)
        assert resp.requires_approval, "Baseline must require approval in adapter"

        bundle = export_bundle(p.root, include_sources=True)
        import zipfile
        with zipfile.ZipFile(bundle) as zf:
            names = zf.namelist()
        assert not any(n.startswith("source/uploads") for n in names), "Bundle must not include original uploads"
        assert any(n.startswith("source/redacted") for n in names), "Bundle should include redacted sources when requested"

        print("NEGATIVE_CHECKS_PASS", root)
    finally:
        shutil.rmtree(root, ignore_errors=True)


if __name__ == "__main__":
    main()
